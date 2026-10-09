from __future__ import annotations

import base64
import json

import httpx
import pytest

from m365_mcp.sharepoint_files import SharePointFilesClient


class StaticAuthService:
    async def get_access_token(self) -> str:
        return "access-token"


def _make_client(handler) -> tuple[SharePointFilesClient, httpx.AsyncClient]:
    http_client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    return SharePointFilesClient(StaticAuthService(), http_client), http_client


@pytest.mark.anyio
async def test_search_sites_builds_query_and_maps_response() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.headers["authorization"] == "Bearer access-token"
        # Item 3 regression guard: no Outlook ImmutableId Prefer header here.
        assert "prefer" not in request.headers
        assert request.url.path == "/v1.0/sites"
        assert request.url.params["search"] == "acme deals"
        assert request.url.params["$top"] == "100"  # clamped from 500
        return httpx.Response(
            200,
            json={
                "value": [
                    {
                        "id": "site-1",
                        "name": "acme",
                        "displayName": "Acme Deals",
                        "webUrl": "https://contoso.sharepoint.com/sites/acme",
                    }
                ]
            },
        )

    client, http_client = _make_client(handler)
    result = await client.search_sites(query="acme deals", top=500)
    assert result.query == "acme deals"
    assert result.sites[0].id == "site-1"
    assert result.sites[0].displayName == "Acme Deals"
    await http_client.aclose()


@pytest.mark.anyio
async def test_get_site_by_path_constructs_colon_path() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert (
            request.url.path
            == "/v1.0/sites/contoso.sharepoint.com:/sites/Acquisitions"
        )
        return httpx.Response(200, json={"id": "site-9", "displayName": "Acq"})

    client, http_client = _make_client(handler)
    site = await client.get_site_by_path(
        hostname="contoso.sharepoint.com", sitePath="/Acquisitions/"
    )
    assert site.id == "site-9"
    await http_client.aclose()


@pytest.mark.anyio
async def test_list_drives_maps_libraries() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/v1.0/sites/site-1/drives"
        return httpx.Response(
            200,
            json={
                "value": [
                    {
                        "id": "drive-1",
                        "name": "Documents",
                        "driveType": "documentLibrary",
                        "webUrl": "https://contoso.sharepoint.com/Docs",
                    }
                ]
            },
        )

    client, http_client = _make_client(handler)
    result = await client.list_drives(siteId="site-1")
    assert result.siteId == "site-1"
    assert result.drives[0].driveType == "documentLibrary"
    await http_client.aclose()


@pytest.mark.anyio
async def test_list_children_by_item_maps_folders_and_files() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/v1.0/drives/drive-1/items/item-1/children"
        assert request.url.params["$top"] == "999"  # clamped from 5000
        return httpx.Response(
            200,
            json={
                "value": [
                    {
                        "id": "f1",
                        "name": "Reports",
                        "folder": {"childCount": 3},
                        "parentReference": {"path": "/drive/root:/Reports"},
                    },
                    {
                        "id": "x1",
                        "name": "Budget.xlsx",
                        "file": {"mimeType": "application/vnd.ms-excel"},
                        "size": 4096,
                        "parentReference": {
                            "driveId": "drive-1",
                            "path": "/drive/root:",
                        },
                    },
                ],
                "@odata.nextLink": "https://graph.microsoft.com/next",
            },
        )

    client, http_client = _make_client(handler)
    result = await client.list_children(
        driveId="drive-1", itemId="item-1", top=5000
    )
    assert result.parentItemId == "item-1"
    assert result.nextLink == "https://graph.microsoft.com/next"
    folder, xlsx = result.items
    assert folder.isFolder is True
    assert folder.childCount == 3
    assert xlsx.isFolder is False
    assert xlsx.fileExtension == "xlsx"
    assert xlsx.driveId == "drive-1"
    await http_client.aclose()


@pytest.mark.anyio
async def test_list_children_by_path_uses_root_colon_addressing() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert (
            request.url.path
            == "/v1.0/drives/drive-1/root:/Shared Active Deals/Q1:/children"
        )
        return httpx.Response(200, json={"value": []})

    client, http_client = _make_client(handler)
    result = await client.list_children(
        driveId="drive-1", path="/Shared Active Deals/Q1/"
    )
    assert result.items == []
    await http_client.aclose()


@pytest.mark.anyio
async def test_list_children_extension_and_folder_filters() -> None:
    value = {
        "value": [
            {"id": "f1", "name": "Sub", "folder": {"childCount": 0}},
            {"id": "a1", "name": "a.xlsx", "file": {}},
            {"id": "b1", "name": "b.pdf", "file": {}},
            {"id": "c1", "name": "c.docx", "file": {}},
        ]
    }

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json=value)

    client, http_client = _make_client(handler)

    only_xlsx = await client.list_children(
        driveId="d", itemId="i", extensions=[".XLSX"]
    )
    names = {i.name for i in only_xlsx.items}
    # Folders are always kept; only xlsx files pass the extension filter.
    assert names == {"Sub", "a.xlsx"}

    folders = await client.list_children(driveId="d", itemId="i", foldersOnly=True)
    assert [i.name for i in folders.items] == ["Sub"]
    await http_client.aclose()


@pytest.mark.anyio
async def test_search_in_drive_escapes_single_quotes() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        # OData single-quote escaping: O'Brien -> O''Brien inside search(q='...').
        assert "search(q='O''Brien')" in str(request.url)
        return httpx.Response(200, json={"value": []})

    client, http_client = _make_client(handler)
    result = await client.search_in_drive(driveId="drive-1", query="O'Brien")
    assert result.driveId == "drive-1"
    await http_client.aclose()


@pytest.mark.anyio
async def test_search_items_posts_search_query_and_parses_hits() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.method == "POST"
        assert request.url.path == "/v1.0/search/query"
        body = json.loads(request.content)
        req = body["requests"][0]
        assert req["entityTypes"] == ["driveItem"]
        assert req["query"]["queryString"] == "tracker"
        assert req["size"] == 200  # clamped from 999
        return httpx.Response(
            200,
            json={
                "value": [
                    {
                        "hitsContainers": [
                            {
                                "hits": [
                                    {
                                        "resource": {
                                            "id": "hit-1",
                                            "name": "Tracker.xlsx",
                                            "file": {},
                                            "parentReference": {"driveId": "drive-9"},
                                        }
                                    }
                                ]
                            }
                        ]
                    }
                ]
            },
        )

    client, http_client = _make_client(handler)
    result = await client.search_items(query="tracker", top=999)
    assert result.items[0].itemId == "hit-1"
    assert result.items[0].driveId == "drive-9"
    assert result.items[0].fileExtension == "xlsx"
    await http_client.aclose()


@pytest.mark.anyio
async def test_get_item_by_share_url_encodes_url() -> None:
    share_url = "https://contoso.sharepoint.com/:x:/r/sites/a/Doc.xlsx?web=1"
    expected = "u!" + base64.urlsafe_b64encode(
        share_url.encode("utf-8")
    ).decode("ascii").rstrip("=")

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path == f"/v1.0/shares/{expected}/driveItem"
        return httpx.Response(
            200,
            json={
                "id": "drv-item",
                "name": "Doc.xlsx",
                "file": {},
                "parentReference": {"driveId": "drive-7"},
            },
        )

    client, http_client = _make_client(handler)
    item = await client.get_item_by_share_url(shareUrl=share_url)
    assert item.itemId == "drv-item"
    assert item.driveId == "drive-7"
    await http_client.aclose()


@pytest.mark.anyio
async def test_list_permissions_maps_links_recipients_and_inheritance() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/v1.0/next":
            return httpx.Response(200, json={"value": []})
        assert request.url.path == "/v1.0/drives/drive-1/items/item-1/permissions"
        return httpx.Response(
            200,
            json={
                "value": [
                    {
                        "id": "perm-link",
                        "roles": ["read"],
                        "link": {
                            "type": "view",
                            "scope": "organization",
                            "webUrl": "https://contoso.sharepoint.com/share/link",
                        },
                        "expirationDateTime": "2026-10-01T00:00:00Z",
                    },
                    {
                        "id": "perm-user",
                        "roles": ["write"],
                        "grantedToV2": {
                            "user": {
                                "id": "user-1",
                                "displayName": "Ada Lovelace",
                                "email": "ada@example.com",
                            }
                        },
                        "grantedToIdentitiesV2": [
                            {
                                "group": {
                                    "id": "group-1",
                                    "displayName": "Project Team",
                                }
                            }
                        ],
                        "inheritedFrom": {"id": "parent-1"},
                    },
                ],
                "@odata.nextLink": "https://graph.microsoft.com/v1.0/next",
            },
        )

    client, http_client = _make_client(handler)
    result = await client.list_permissions(driveId="drive-1", itemId="item-1")
    link, user = result.permissions
    assert link.permissionId == "perm-link"
    assert link.linkType == "view"
    assert link.linkScope == "organization"
    assert link.shareUrl == "https://contoso.sharepoint.com/share/link"
    assert user.inherited is True
    assert user.grantedTo[0].email == "ada@example.com"
    assert user.grantedTo[1].identityType == "group"
    assert user.grantedTo[1].displayName == "Project Team"
    assert result.nextLink is None
    await http_client.aclose()


@pytest.mark.anyio
async def test_create_link_preserves_inheritance_and_maps_result() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.method == "POST"
        assert request.url.path == "/v1.0/drives/drive-1/items/item-1/createLink"
        assert json.loads(request.content) == {
            "type": "view",
            "scope": "organization",
            "retainInheritedPermissions": True,
            "expirationDateTime": "2026-10-01T00:00:00Z",
        }
        return httpx.Response(
            201,
            json={
                "id": "perm-1",
                "roles": ["read"],
                "link": {
                    "type": "view",
                    "scope": "organization",
                    "webUrl": "https://contoso.sharepoint.com/share/link",
                },
            },
        )

    client, http_client = _make_client(handler)
    result = await client.create_link(
        driveId="drive-1",
        itemId="item-1",
        linkType="view",
        scope="organization",
        expirationDateTime="2026-10-01T00:00:00Z",
    )
    assert result.permission.permissionId == "perm-1"
    assert result.permission.roles == ["read"]
    await http_client.aclose()


@pytest.mark.anyio
async def test_grant_access_posts_invite_without_email_by_default() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/v1.0/drives/drive-1/items/item-1/invite"
        assert json.loads(request.content) == {
            "recipients": [{"email": "ada@example.com"}],
            "roles": ["read"],
            "requireSignIn": True,
            "sendInvitation": False,
            "retainInheritedPermissions": True,
        }
        return httpx.Response(
            200,
            json={
                "value": [
                    {
                        "id": "perm-user",
                        "roles": ["read"],
                        "invitation": {"email": "ada@example.com"},
                    }
                ]
            },
        )

    client, http_client = _make_client(handler)
    result = await client.grant_access(
        driveId="drive-1",
        itemId="item-1",
        recipients=["  ada@example.com  "],
        role="read",
    )
    assert result.permissions[0].grantedTo[0].email == "ada@example.com"
    assert result.failures == []
    await http_client.aclose()


@pytest.mark.anyio
async def test_grant_access_reports_partial_failures() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            207,
            json={
                "value": [
                    {
                        "id": "perm-user",
                        "roles": ["read"],
                        "invitation": {"email": "ada@example.com"},
                    },
                    {
                        "id": "perm-external",
                        "roles": ["read"],
                        "invitation": {"email": "external@example.net"},
                        "error": {
                            "code": "notAllowed",
                            "message": "Access granted, but notification failed.",
                        }
                    },
                ]
            },
        )

    client, http_client = _make_client(handler)
    result = await client.grant_access(
        driveId="drive-1",
        itemId="item-1",
        recipients=["ada@example.com", "external@example.net"],
        role="read",
    )
    assert len(result.permissions) == 2
    assert result.permissions[1].permissionId == "perm-external"
    assert result.failures[0].recipient == "external@example.net"
    assert result.failures[0].code == "notAllowed"
    await http_client.aclose()


@pytest.mark.anyio
async def test_revoke_permission_deletes_only_the_permission() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.method == "DELETE"
        assert (
            request.url.path
            == "/v1.0/drives/drive-1/items/item-1/permissions/perm-1"
        )
        return httpx.Response(204)

    client, http_client = _make_client(handler)
    result = await client.revoke_permission(
        driveId="drive-1", itemId="item-1", permissionId="perm-1"
    )
    assert result.revoked is True
    await http_client.aclose()


@pytest.mark.anyio
async def test_grant_access_requires_a_recipient() -> None:
    client, http_client = _make_client(
        lambda request: pytest.fail(f"unexpected request: {request.url}")
    )
    with pytest.raises(ValueError, match="recipient"):
        await client.grant_access(
            driveId="drive-1", itemId="item-1", recipients=["  "], role="read"
        )
    await http_client.aclose()


@pytest.mark.anyio
async def test_request_error_raises_with_graph_detail() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            403,
            json={
                "error": {
                    "code": "accessDenied",
                    "message": "Access denied to the resource.",
                }
            },
        )

    client, http_client = _make_client(handler)
    with pytest.raises(RuntimeError) as excinfo:
        await client.list_drives(siteId="site-1")
    message = str(excinfo.value)
    assert "403" in message
    assert "accessDenied" in message
    assert "Access denied to the resource." in message
    await http_client.aclose()


def test_encode_share_url_matches_graph_addressing() -> None:
    encoded = SharePointFilesClient._encode_share_url("https://a/b c?d=1")
    assert encoded.startswith("u!")
    assert "=" not in encoded  # padding stripped
    # Round-trips back to the original URL.
    payload = encoded[2:]
    padded = payload + "=" * (-len(payload) % 4)
    assert base64.urlsafe_b64decode(padded).decode("utf-8") == "https://a/b c?d=1"


def test_q_escapes_single_quote() -> None:
    assert SharePointFilesClient._q("it's a 'test'") == "it''s a ''test''"


# --------------------------------------------------------------------------- #
# get_file_content
# --------------------------------------------------------------------------- #
DOWNLOAD_URL = "https://contoso-my.sharepoint.com/download.aspx?tempauth=abc"


def _docx_bytes() -> bytes:
    import io

    from docx import Document

    document = Document()
    document.add_heading("Investment Memo", level=0)
    document.add_heading("Summary", level=1)
    document.add_paragraph("Acme Apartments is a 240-unit property.")
    document.add_paragraph("Strong submarket", style="List Bullet")
    document.add_paragraph("Value-add upside", style="List Bullet")
    document.add_heading("Returns", level=2)
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("Metric", "Value"), ("IRR", "15.2%"), ("Note", "A | B")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("Closing remarks.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _file_handler(name: str, content: bytes, *, size: int | None = None, seen: list | None = None):
    def handler(request: httpx.Request) -> httpx.Response:
        if seen is not None:
            seen.append(request)
        if request.url.host == "graph.microsoft.com":
            assert request.url.path == "/v1.0/drives/drive-1/items/item-1"
            assert "@microsoft.graph.downloadUrl" in request.url.params["$select"]
            return httpx.Response(
                200,
                json={
                    "id": "item-1",
                    "name": name,
                    "file": {"mimeType": "application/octet-stream"},
                    "size": len(content) if size is None else size,
                    "parentReference": {"driveId": "drive-1", "path": "/drive/root:/Deals"},
                    "@microsoft.graph.downloadUrl": DOWNLOAD_URL,
                },
            )
        assert str(request.url) == DOWNLOAD_URL
        return httpx.Response(200, content=content)

    return handler


@pytest.mark.anyio
async def test_get_file_content_reads_docx_as_markdown_without_token_on_download() -> None:
    seen: list[httpx.Request] = []
    client, http_client = _make_client(_file_handler("Memo.docx", _docx_bytes(), seen=seen))
    result = await client.get_file_content(driveId="drive-1", itemId="item-1")

    assert result.unsupportedReason is None
    assert result.encoding == "docx-markdown"
    assert result.item.name == "Memo.docx"
    assert result.content == (
        "# Investment Memo\n\n"
        "# Summary\n\n"
        "Acme Apartments is a 240-unit property.\n\n"
        "- Strong submarket\n"
        "- Value-add upside\n\n"
        "## Returns\n\n"
        "| Metric | Value |\n"
        "| --- | --- |\n"
        "| IRR | 15.2% |\n"
        "| Note | A \\| B |\n\n"
        "Closing remarks."
    )
    # The pre-authenticated download URL must never receive the bearer token.
    assert seen[0].headers["authorization"] == "Bearer access-token"
    assert "authorization" not in seen[1].headers
    await http_client.aclose()


@pytest.mark.anyio
async def test_get_file_content_reads_markdown_and_truncates_to_max_chars() -> None:
    text = "﻿# Notes\n\nCafé — deal notes.\n".encode("utf-8")
    client, http_client = _make_client(_file_handler("notes.md", text))
    result = await client.get_file_content(driveId="drive-1", itemId="item-1", maxChars=7)

    assert result.encoding == "utf-8"
    assert result.content == "# Notes"
    assert result.truncated is True
    assert result.unsupportedReason is None
    await http_client.aclose()


@pytest.mark.anyio
async def test_get_file_content_extracts_pdf_text() -> None:
    from tests.test_graph import _single_page_pdf

    client, http_client = _make_client(
        _file_handler("OM.pdf", _single_page_pdf(text=b"RENT ROLL", pages=2))
    )
    result = await client.get_file_content(driveId="drive-1", itemId="item-1")

    assert result.encoding == "pdf-text"
    assert "--- Page 1 ---\nRENT ROLL p1" in result.content
    assert "--- Page 2 ---\nRENT ROLL p2" in result.content
    await http_client.aclose()


@pytest.mark.anyio
@pytest.mark.parametrize(
    ("name", "reason_fragment"),
    [
        ("Model.xlsx", "workbook tools"),
        ("Old.doc", "save the document as .docx"),
        ("photo.png", ".png files are not readable"),
    ],
)
async def test_get_file_content_refuses_unreadable_types_before_download(
    name: str, reason_fragment: str
) -> None:
    seen: list[httpx.Request] = []
    client, http_client = _make_client(_file_handler(name, b"bytes", seen=seen))
    result = await client.get_file_content(driveId="drive-1", itemId="item-1")

    assert result.content is None
    assert reason_fragment in result.unsupportedReason
    assert len(seen) == 1  # metadata only; nothing downloaded
    await http_client.aclose()


@pytest.mark.anyio
async def test_get_file_content_refuses_folders() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"id": "item-1", "name": "Deals", "folder": {"childCount": 3}})

    client, http_client = _make_client(handler)
    result = await client.get_file_content(driveId="drive-1", itemId="item-1")
    assert "sharepoint_list_children" in result.unsupportedReason
    await http_client.aclose()


@pytest.mark.anyio
async def test_get_file_content_enforces_max_bytes_on_metadata_and_stream() -> None:
    seen: list[httpx.Request] = []
    client, http_client = _make_client(
        _file_handler("big.txt", b"x" * 50, size=50, seen=seen)
    )
    by_size = await client.get_file_content(driveId="drive-1", itemId="item-1", maxBytes=10)
    assert by_size.unsupportedReason == "File size exceeds maxBytes=10"
    assert len(seen) == 1

    # Metadata under-reports the size: the download stops once it passes maxBytes.
    client, http_client2 = _make_client(_file_handler("big.txt", b"x" * 50, size=5))
    by_stream = await client.get_file_content(driveId="drive-1", itemId="item-1", maxBytes=10)
    assert by_stream.truncated is True
    assert by_stream.content is None
    assert by_stream.unsupportedReason == "File content exceeds maxBytes=10"
    await http_client.aclose()
    await http_client2.aclose()


@pytest.mark.anyio
async def test_get_file_content_follows_content_redirect_without_token() -> None:
    seen: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request)
        if request.url.path == "/v1.0/drives/drive-1/items/item-1":
            return httpx.Response(200, json={"id": "item-1", "name": "readme.txt", "file": {}, "size": 5})
        if request.url.path == "/v1.0/drives/drive-1/items/item-1/content":
            assert request.headers["authorization"] == "Bearer access-token"
            return httpx.Response(302, headers={"location": DOWNLOAD_URL})
        assert str(request.url) == DOWNLOAD_URL
        assert "authorization" not in request.headers
        return httpx.Response(200, content=b"hello")

    client, http_client = _make_client(handler)
    result = await client.get_file_content(driveId="drive-1", itemId="item-1")
    assert result.content == "hello"
    assert len(seen) == 3
    await http_client.aclose()


@pytest.mark.anyio
@pytest.mark.parametrize(
    ("name", "content", "reason_fragment"),
    [
        ("Memo.docx", b"\xd0\xcf\x11\xe0 not a zip", "not a valid .docx"),
        ("Scan.pdf", None, "scanned PDFs need OCR"),
    ],
)
async def test_get_file_content_reports_unreadable_documents(
    name: str, content: bytes | None, reason_fragment: str
) -> None:
    if content is None:
        from tests.test_graph import _single_page_pdf

        content = _single_page_pdf(text=b"", pages=1).replace(b"( p1) Tj", b"() Tj")
    client, http_client = _make_client(_file_handler(name, content))
    result = await client.get_file_content(driveId="drive-1", itemId="item-1")
    assert result.content is None
    assert reason_fragment in result.unsupportedReason
    await http_client.aclose()


@pytest.mark.anyio
async def test_get_file_content_reports_download_errors() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.host == "graph.microsoft.com":
            return httpx.Response(
                200,
                json={"id": "item-1", "name": "a.txt", "file": {}, "size": 1,
                      "@microsoft.graph.downloadUrl": DOWNLOAD_URL},
            )
        return httpx.Response(403, json={"error": {"code": "accessDenied", "message": "nope"}})

    client, http_client = _make_client(handler)
    with pytest.raises(RuntimeError, match=r"download failed \(403\): accessDenied: nope"):
        await client.get_file_content(driveId="drive-1", itemId="item-1")
    await http_client.aclose()


def _relabel_main_part(content: bytes, content_type: str) -> bytes:
    """Save a .docx as another Word package type by swapping its main content type."""
    import io
    import zipfile

    docx_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(
        output, "w", zipfile.ZIP_DEFLATED
    ) as copy:
        for info in source.infolist():
            data = source.read(info)
            if info.filename == "[Content_Types].xml":
                assert docx_type.encode() in data
                data = data.replace(docx_type.encode(), content_type.encode())
            copy.writestr(info.filename, data)
    return output.getvalue()


@pytest.mark.anyio
@pytest.mark.parametrize(
    ("name", "content_type"),
    [
        ("Memo.docm", "application/vnd.ms-word.document.macroEnabled.main+xml"),
        (
            "Memo.dotx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
        ),
        ("Memo.dotm", "application/vnd.ms-word.template.macroEnabledTemplate.main+xml"),
    ],
)
async def test_get_file_content_reads_macro_enabled_and_template_word_files(
    name: str, content_type: str
) -> None:
    content = _relabel_main_part(_docx_bytes(), content_type)
    client, http_client = _make_client(_file_handler(name, content))
    result = await client.get_file_content(driveId="drive-1", itemId="item-1")

    assert result.unsupportedReason is None
    assert result.encoding == "docx-markdown"
    assert result.content.startswith("# Investment Memo\n\n# Summary")
    await http_client.aclose()


@pytest.mark.anyio
async def test_get_file_content_clamps_max_bytes_and_max_chars() -> None:
    from m365_mcp.sharepoint_files import MAX_FILE_MAX_BYTES, MAX_FILE_MAX_CHARS

    seen: list[httpx.Request] = []
    client, http_client = _make_client(
        _file_handler("huge.txt", b"x", size=MAX_FILE_MAX_BYTES + 1, seen=seen)
    )
    by_size = await client.get_file_content(
        driveId="drive-1", itemId="item-1", maxBytes=10**12
    )
    assert by_size.unsupportedReason == f"File size exceeds maxBytes={MAX_FILE_MAX_BYTES}"
    assert len(seen) == 1

    text = b"y" * (MAX_FILE_MAX_CHARS + 10)
    client, http_client2 = _make_client(_file_handler("long.txt", text))
    by_chars = await client.get_file_content(
        driveId="drive-1", itemId="item-1", maxChars=10**9
    )
    assert len(by_chars.content) == MAX_FILE_MAX_CHARS
    assert by_chars.truncated is True
    await http_client.aclose()
    await http_client2.aclose()


@pytest.mark.anyio
async def test_get_file_content_refuses_oversized_word_archive() -> None:
    import io
    import zipfile

    from m365_mcp.document_text import MAX_DOCX_INFLATED_BYTES

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\0" * (MAX_DOCX_INFLATED_BYTES + 1))
    client, http_client = _make_client(_file_handler("Bomb.docx", buffer.getvalue()))
    result = await client.get_file_content(driveId="drive-1", itemId="item-1")

    assert result.content is None
    assert f"over the {MAX_DOCX_INFLATED_BYTES}-byte limit" in result.unsupportedReason
    await http_client.aclose()


@pytest.mark.anyio
@pytest.mark.parametrize("via_redirect", [False, True])
async def test_get_file_content_refuses_non_https_download_urls(via_redirect: bool) -> None:
    insecure_url = "http://contoso-my.sharepoint.com/download.aspx?tempauth=abc"
    downloads: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/v1.0/drives/drive-1/items/item-1":
            data = {"id": "item-1", "name": "a.txt", "file": {}, "size": 1}
            if not via_redirect:
                data["@microsoft.graph.downloadUrl"] = insecure_url
            return httpx.Response(200, json=data)
        if request.url.path == "/v1.0/drives/drive-1/items/item-1/content":
            return httpx.Response(302, headers={"location": insecure_url})
        downloads.append(request)
        return httpx.Response(200, content=b"a")

    client, http_client = _make_client(handler)
    with pytest.raises(RuntimeError, match="not https"):
        await client.get_file_content(driveId="drive-1", itemId="item-1")
    assert downloads == []
    await http_client.aclose()


@pytest.mark.anyio
async def test_get_file_content_parses_off_the_event_loop() -> None:
    import threading

    from m365_mcp import sharepoint_files

    loop_thread = threading.get_ident()
    parse_threads: list[int] = []

    def decode(content: bytes) -> str:
        parse_threads.append(threading.get_ident())
        return content.decode()

    client, http_client = _make_client(_file_handler("a.txt", b"hello"))
    original = sharepoint_files.decode_text
    sharepoint_files.decode_text = decode
    try:
        result = await client.get_file_content(driveId="drive-1", itemId="item-1")
    finally:
        sharepoint_files.decode_text = original
    assert result.content == "hello"
    assert parse_threads and parse_threads[0] != loop_thread
    await http_client.aclose()
