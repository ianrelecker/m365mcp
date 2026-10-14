from __future__ import annotations

import base64
import json

import httpx
import pytest

from m365_mcp import microsoft_graph as graph_module
from m365_mcp import workbook_reader
from m365_mcp.microsoft_graph import MicrosoftGraphClient


class StaticAuthService:
    async def get_access_token(self) -> str:
        return "access-token"


@pytest.mark.anyio
async def test_list_messages_and_search_messages() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.headers["authorization"] == "Bearer access-token"

        if request.url.path.endswith("/mailFolders('Inbox')/messages"):
            assert request.url.params["$top"] == "100"
            assert "inferenceClassification" in request.url.params["$select"]
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "m1",
                            "subject": "Hello",
                            "from": {"emailAddress": {"address": "sender@example.com"}},
                            "receivedDateTime": "2026-04-21T12:00:00Z",
                            "sentDateTime": "2026-04-21T12:00:00Z",
                            "bodyPreview": "Preview",
                            "webLink": "https://outlook.example/messages/m1",
                            "isDraft": False,
                            "inferenceClassification": "focused",
                            "conversationId": "conv-1",
                        }
                    ]
                },
            )

        if request.url.path.endswith("/messages"):
            assert request.headers["consistencylevel"] == "eventual"
            assert request.url.params["$top"] == "50"
            assert request.url.params["$search"] == "\"from:\\\"boss\\\"\""
            assert "inferenceClassification" in request.url.params["$select"]
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "m2",
                            "subject": "Search hit",
                            "bodyPreview": "",
                            "inferenceClassification": "other",
                        }
                    ]
                },
            )

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    listed = await graph.list_messages(mailbox=None, folder="Inbox", top=999)
    assert listed.mailbox == "me"
    assert listed.messages[0].from_ == "sender@example.com"
    assert listed.messages[0].inferenceClassification == "focused"

    searched = await graph.search_messages(query='from:"boss"', top=99)
    assert searched.mailbox == "me"
    assert searched.messages[0].id == "m2"
    assert searched.messages[0].inferenceClassification == "other"

    await client.aclose()


@pytest.mark.anyio
async def test_list_messages_filters_by_date_sorts_and_pages() -> None:
    requests: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        skip = int(request.url.params.get("$skip", "0"))
        body: dict[str, object] = {
            "value": [
                {"id": f"m{skip + index}", "subject": "", "bodyPreview": "", "isDraft": False}
                for index in range(2)
            ]
        }
        if skip == 0:
            # With a filter, Graph's continuation $skip counts items scanned,
            # not messages returned, so it differs from len(value).
            body["@odata.nextLink"] = (
                "https://graph.microsoft.com/v1.0/me/mailFolders('Inbox')/messages"
                "?$top=2&$filter=x&$skip=7"
            )
        return httpx.Response(200, json=body)

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    first = await graph.list_messages(
        top=2,
        inferenceClassification="focused",
        receivedAfter="2026-10-01",
        receivedBefore="2026-10-07T09:00:00-05:00",
    )
    params = requests[-1].url.params
    # The sorted property leads the filter, as Graph requires for $orderby.
    assert params["$filter"] == (
        "receivedDateTime ge 2026-10-01T00:00:00Z and "
        "receivedDateTime lt 2026-10-07T14:00:00Z and "
        "inferenceClassification eq 'focused'"
    )
    assert params["$orderby"] == "receivedDateTime desc"
    assert "$skip" not in params
    assert first.receivedAfter == "2026-10-01T00:00:00Z"
    assert first.receivedBefore == "2026-10-07T14:00:00Z"
    assert first.hasMore is True
    assert first.nextCursor

    second = await graph.list_messages(
        top=2,
        inferenceClassification="focused",
        receivedAfter="2026-10-01",
        receivedBefore="2026-10-07T09:00:00-05:00",
        cursor=first.nextCursor,
    )
    params = requests[-1].url.params
    # Graph's own continuation, not skip + len(messages); the filter comes
    # from the arguments, not from the link.
    assert params["$skip"] == "7"
    assert params["$filter"].startswith("receivedDateTime ge 2026-10-01T00:00:00Z")
    assert [message.id for message in second.messages] == ["m7", "m8"]
    assert second.hasMore is False
    assert second.nextCursor is None

    # A cursor is tied to the query that produced it: changing the filter,
    # order, or folder would make Graph's scanned-item $skip skip or repeat.
    requests_before = len(requests)
    for changed in (
        {"inferenceClassification": "other"},
        {"order": "oldest"},
        {"folder": "Archive"},
        {"mailbox": "shared@example.com"},
        {"receivedAfter": "2026-10-02"},
    ):
        arguments: dict[str, object] = {
            "top": 2,
            "inferenceClassification": "focused",
            "receivedAfter": "2026-10-01",
            "receivedBefore": "2026-10-07T09:00:00-05:00",
            **changed,
        }
        with pytest.raises(ValueError, match="different mail_list arguments"):
            await graph.list_messages(cursor=first.nextCursor, **arguments)
    assert len(requests) == requests_before
    # Page size is not part of the query identity.
    await graph.list_messages(
        top=5,
        inferenceClassification="focused",
        receivedAfter="2026-10-01",
        receivedBefore="2026-10-07T09:00:00-05:00",
        cursor=first.nextCursor,
    )
    assert requests[-1].url.params["$skip"] == "7"

    await graph.list_messages(top=2, order="oldest")
    params = requests[-1].url.params
    assert params["$orderby"] == "receivedDateTime asc"
    # No filters at all: plain $orderby needs no date bound.
    assert "$filter" not in params

    await graph.list_messages(receivedBefore="2026-10-07Z")
    assert requests[-1].url.params["$filter"] == (
        "receivedDateTime ge 1900-01-01T00:00:00Z and "
        "receivedDateTime lt 2026-10-07T00:00:00Z"
    )

    with pytest.raises(ValueError, match="receivedAfter must be an ISO 8601"):
        await graph.list_messages(receivedAfter="last week")
    with pytest.raises(ValueError, match="earlier than"):
        await graph.list_messages(receivedAfter="2026-10-07", receivedBefore="2026-10-01")
    with pytest.raises(ValueError, match="order"):
        await graph.list_messages(order="latest")
    for bad_cursor in (
        "not-a-cursor",
        base64.urlsafe_b64encode(b"$filter=x").decode(),
        # No scope fingerprint.
        base64.urlsafe_b64encode(b"$skip=7").decode(),
        # Non-ASCII digits pass str.isdigit() but not Graph.
        base64.urlsafe_b64encode("$skip=\u00b2&scope=abc".encode()).decode(),
        base64.urlsafe_b64encode(b"$skip=7&$skip=8&scope=abc").decode(),
    ):
        with pytest.raises(ValueError, match="cursor"):
            await graph.list_messages(cursor=bad_cursor)

    await client.aclose()


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("2026-10-01", "2026-10-01T00:00:00Z"),
        ("2026-10-07Z", "2026-10-07T00:00:00Z"),
        ("2026-10-01T12:00:00.500Z", "2026-10-01T12:00:00.500Z"),
        ("2026-10-01T12:00:00.9Z", "2026-10-01T12:00:00.9Z"),
        ("2026-10-01T12:00:00.1234567Z", "2026-10-01T12:00:00.1234567Z"),
        ("2026-10-01T09:00:00.25-05:00", "2026-10-01T14:00:00.25Z"),
        # Formats datetime.fromisoformat only accepts on Python 3.11+.
        ("2026-10-01T09:00-0500", "2026-10-01T14:00:00Z"),
        ("2026-10-01T09:00:00+05", "2026-10-01T04:00:00Z"),
        ("2026-10-01 09:00:00,5", "2026-10-01T09:00:00.5Z"),
        ("2026-10-01T23:30:00-01:00", "2026-10-02T00:30:00Z"),
    ],
)
def test_graph_datetime_normalizes_to_utc(value: str, expected: str) -> None:
    result = MicrosoftGraphClient._graph_datetime(value, "receivedAfter")
    assert result is not None
    assert result[0] == expected


@pytest.mark.parametrize(
    ("value", "message"),
    [
        ("last week", "ISO 8601"),
        ("20261001", "ISO 8601"),
        ("2026-13-01", "ISO 8601"),
        ("2026-10-01T25:00:00Z", "ISO 8601"),
        ("2026-10-01T09:00:00+24:00", "ISO 8601"),
        ("2026-10-01T09:00:00+05:99", "ISO 8601"),
        # Graph's DateTimeOffset holds at most 7 fractional digits.
        ("2026-10-01T09:00:00.12345678Z", "ISO 8601"),
        ("9999-12-31T23:00:00-05:00", "ISO 8601"),
        ("0999-01-01", "1900 or later"),
        ("1899-12-31", "1900 or later"),
    ],
)
def test_graph_datetime_rejects_invalid_values(value: str, message: str) -> None:
    with pytest.raises(ValueError, match=message):
        MicrosoftGraphClient._graph_datetime(value, "receivedAfter")


@pytest.mark.anyio
async def test_list_messages_compares_fractional_bounds_numerically() -> None:
    client = httpx.AsyncClient(
        transport=httpx.MockTransport(lambda request: httpx.Response(200, json={"value": []}))
    )
    graph = MicrosoftGraphClient(StaticAuthService(), client)
    # As strings ".5Z" sorts before "Z"; as times, .5 s is later.
    result = await graph.list_messages(
        receivedAfter="2026-10-01T12:00:00Z",
        receivedBefore="2026-10-01T12:00:00.5Z",
    )
    assert result.receivedBefore == "2026-10-01T12:00:00.5Z"
    with pytest.raises(ValueError, match="earlier than"):
        await graph.list_messages(
            receivedAfter="2026-10-01T12:00:00.5Z",
            receivedBefore="2026-10-01T12:00:00Z",
        )
    await client.aclose()


def test_mail_list_cursor_only_from_graph_links() -> None:
    def next_cursor(link: str | None) -> str | None:
        return MicrosoftGraphClient._mail_list_next_cursor(link, "scope1")

    assert next_cursor(None) is None
    assert next_cursor("https://example.com/v1.0/me/messages?$skip=7") is None
    assert next_cursor("http://graph.microsoft.com/v1.0/me/messages?$skip=7") is None
    assert next_cursor("https://graph.microsoft.com/v1.0/me/messages?$top=5") is None
    cursor = next_cursor(
        "https://graph.microsoft.com/v1.0/me/messages?$top=5&$skiptoken=abc%3D"
    )
    assert MicrosoftGraphClient._mail_list_cursor_params(cursor) == (
        {"$skiptoken": "abc="},
        "scope1",
    )


@pytest.mark.anyio
async def test_mail_list_cursor_values_stay_inside_one_parameter() -> None:
    requests: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        body: dict[str, object] = {"value": []}
        if "$skiptoken" not in request.url.params:
            body["@odata.nextLink"] = (
                "https://graph.microsoft.com/v1.0/me/mailFolders('Inbox')/messages"
                "?$skiptoken=a%26%24filter%3Dx%23y"
            )
        return httpx.Response(200, json=body)

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)
    first = await graph.list_messages(isRead=False)
    await graph.list_messages(isRead=False, cursor=first.nextCursor)
    params = requests[-1].url.params
    assert params["$skiptoken"] == "a&$filter=x#y"
    assert params["$filter"] == (
        "receivedDateTime ge 1900-01-01T00:00:00Z and isRead eq false"
    )
    await client.aclose()


@pytest.mark.anyio
async def test_mail_list_has_more_only_with_a_usable_cursor() -> None:
    client = httpx.AsyncClient(
        transport=httpx.MockTransport(
            lambda request: httpx.Response(
                200,
                json={
                    "value": [],
                    "@odata.nextLink": "https://example.com/v1.0/me/messages?$skip=5",
                },
            )
        )
    )
    graph = MicrosoftGraphClient(StaticAuthService(), client)
    result = await graph.list_messages()
    assert result.nextCursor is None
    assert result.hasMore is False
    await client.aclose()


@pytest.mark.anyio
async def test_get_message_and_list_drafts() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/messages/msg-123"):
            assert "inferenceClassification" in request.url.params["$select"]
            return httpx.Response(
                200,
                json={
                    "id": "msg-123",
                    "subject": "Full message",
                    "from": {"emailAddress": {"address": "sender@example.com"}},
                    "toRecipients": [{"emailAddress": {"address": "to@example.com"}}],
                    "ccRecipients": [{"emailAddress": {"address": "cc@example.com"}}],
                    "bccRecipients": [{"emailAddress": {"address": "bcc@example.com"}}],
                    "bodyPreview": "Preview",
                    "body": {"contentType": "html", "content": "<p>Hello</p>"},
                    "isDraft": False,
                    "inferenceClassification": "focused",
                },
            )

        if request.url.path.endswith("/mailFolders('Drafts')/messages"):
            return httpx.Response(
                200,
                json={"value": [{"id": "draft-1", "subject": "Draft", "isDraft": True}]},
            )

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    message = await graph.get_message(mailbox="shared@example.com", messageId="msg-123")
    assert message.mailbox == "shared@example.com"
    assert message.message.to == ["to@example.com"]
    assert message.message.body.contentType == "html"
    assert message.message.inferenceClassification == "focused"

    drafts = await graph.list_drafts(mailbox="shared@example.com", top=10)
    assert drafts.drafts[0].isDraft is True

    await client.aclose()


@pytest.mark.anyio
async def test_get_message_returns_plain_text_new_part_by_default() -> None:
    requests: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        message_id = request.url.path.rsplit("/", 1)[-1]
        prefers_text = 'outlook.body-content-type="text"' in request.headers["prefer"]
        unique = {
            "msg-reply": "Sounds good, see you Tuesday.",
            "msg-empty": "",
        }[message_id]
        return httpx.Response(
            200,
            json={
                "id": message_id,
                "subject": "Re: Tour",
                "bodyPreview": "Sounds good",
                "isDraft": False,
                "body": {
                    "contentType": "text" if prefers_text else "html",
                    "content": "Sounds good, see you Tuesday.\n\nOn Monday Bob wrote: ...",
                },
                "uniqueBody": {
                    "contentType": "text" if prefers_text else "html",
                    "content": unique
                    if prefers_text
                    else f"<html><head><style>p {{}}</style></head><body>{unique}</body></html>",
                },
            },
        )

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    reply = await graph.get_message(messageId="msg-reply")
    assert reply.message.bodyScope == "new"
    assert reply.message.body.contentType == "text"
    assert reply.message.body.content == "Sounds good, see you Tuesday."
    prefer = requests[-1].headers["prefer"]
    assert 'IdType="ImmutableId"' in prefer
    assert 'outlook.body-content-type="text"' in prefer
    assert "uniqueBody" in requests[-1].url.params["$select"].split(",")

    # An empty uniqueBody (here an HTML shell with only a style block) falls
    # back to the full body rather than returning nothing.
    empty = await graph.get_message(messageId="msg-empty", bodyFormat="html")
    assert 'outlook.body-content-type="html"' in requests[-1].headers["prefer"]
    assert empty.message.bodyScope == "full"
    assert empty.message.body.content.startswith("Sounds good")

    full = await graph.get_message(messageId="msg-reply", bodyScope="full")
    assert "uniqueBody" not in requests[-1].url.params["$select"].split(",")
    assert full.message.bodyScope == "full"
    assert "Bob wrote" in full.message.body.content

    with pytest.raises(ValueError, match="bodyFormat"):
        await graph.get_message(messageId="msg-reply", bodyFormat="markdown")
    with pytest.raises(ValueError, match="bodyScope"):
        await graph.get_message(messageId="msg-reply", bodyScope="quoted")

    await client.aclose()


@pytest.mark.anyio
async def test_get_messages_batches_reads_and_keeps_order() -> None:
    batches: list[dict[str, object]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.method == "POST"
        assert request.url.path == "/v1.0/$batch"
        payload = json.loads(request.content)
        batches.append(payload)
        responses = []
        for inner in payload["requests"]:
            assert inner["method"] == "GET"
            assert inner["url"].startswith("/users/shared%40example.com/messages/")
            assert "uniqueBody" in inner["url"]
            assert 'outlook.body-content-type="text"' in inner["headers"]["Prefer"]
            message_id = inner["url"].split("/messages/", 1)[1].split("?", 1)[0]
            if message_id == "missing":
                responses.append(
                    {
                        "id": inner["id"],
                        "status": 404,
                        "body": {
                            "error": {
                                "code": "ErrorItemNotFound",
                                "message": "The specified object was not found.",
                            }
                        },
                    }
                )
                continue
            responses.append(
                {
                    "id": inner["id"],
                    "status": 200,
                    "body": {
                        "id": message_id,
                        "subject": f"Subject {message_id}",
                        "bodyPreview": "",
                        "isDraft": False,
                        "body": {"contentType": "text", "content": "full"},
                        "uniqueBody": {"contentType": "text", "content": f"new {message_id}"},
                    },
                }
            )
        # Graph does not promise responses in request order.
        return httpx.Response(200, json={"responses": list(reversed(responses))})

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    ids = [f"m{index}" for index in range(22)]
    ids[5] = "missing"
    result = await graph.get_messages(mailbox="shared@example.com", messageIds=ids)

    assert [len(batch["requests"]) for batch in batches] == [20, 2]
    assert result.mailbox == "shared@example.com"
    assert [item.messageId for item in result.messages] == ids
    assert result.messages[0].message.body.content == "new m0"
    assert result.messages[0].message.bodyScope == "new"
    assert result.messages[21].message.id == "m21"
    assert result.messages[5].message is None
    assert "404" in result.messages[5].error
    assert "ErrorItemNotFound" in result.messages[5].error

    with pytest.raises(ValueError, match="at most 50"):
        await graph.get_messages(messageIds=[f"m{index}" for index in range(51)])
    with pytest.raises(ValueError, match="at least one"):
        await graph.get_messages(messageIds=[])

    await client.aclose()


@pytest.mark.anyio
async def test_get_messages_retries_throttled_items(monkeypatch: pytest.MonkeyPatch) -> None:
    sleeps: list[float] = []

    async def fake_sleep(seconds: float) -> None:
        sleeps.append(seconds)

    monkeypatch.setattr(graph_module.anyio, "sleep", fake_sleep)
    attempts: dict[str, int] = {}
    batch_sizes: list[int] = []

    def handler(request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        batch_sizes.append(len(payload["requests"]))
        responses = []
        for inner in payload["requests"]:
            message_id = inner["url"].split("/messages/", 1)[1].split("?", 1)[0]
            attempts[message_id] = attempts.get(message_id, 0) + 1
            # m1 is throttled once, m2 always; m0 succeeds first time.
            if (message_id == "m1" and attempts[message_id] == 1) or message_id == "m2":
                responses.append(
                    {
                        "id": inner["id"],
                        "status": 429,
                        "headers": {"Retry-After": "5"},
                        "body": {"error": {"code": "TooManyRequests", "message": "slow down"}},
                    }
                )
                continue
            responses.append(
                {
                    "id": inner["id"],
                    "status": 200,
                    "body": {
                        "id": message_id,
                        "subject": "",
                        "bodyPreview": "",
                        "isDraft": False,
                        "body": {"contentType": "text", "content": message_id},
                    },
                }
            )
        return httpx.Response(200, json={"responses": responses})

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    result = await graph.get_messages(messageIds=["m0", "m1", "m2"])

    # Only the throttled items are resent, and Retry-After is honoured.
    assert batch_sizes == [3, 2, 1, 1]
    assert sleeps == [5.0, 5.0, 5.0]
    assert attempts == {"m0": 1, "m1": 2, "m2": 4}
    assert result.messages[0].message.id == "m0"
    assert result.messages[1].message.id == "m1"
    assert result.messages[2].message is None
    assert "429" in result.messages[2].error

    await client.aclose()


def test_batch_retry_after_is_bounded() -> None:
    wait = MicrosoftGraphClient._retry_after_seconds
    assert wait({"headers": {"retry-after": "3"}}, 0) == 3.0
    assert wait({"headers": {"Retry-After": "3600"}}, 0) == 20.0
    # Missing or unparseable Retry-After backs off exponentially.
    assert wait({}, 0) == 1.0
    assert wait({"headers": {"Retry-After": "soon"}}, 2) == 4.0
    # HTTP-date form is not parsed; fall back to backoff.
    assert wait({"headers": {"Retry-After": "Wed, 21 Oct 2026 07:28:00 GMT"}}, 1) == 2.0
    assert wait({"headers": {"Retry-After": "nan"}}, 0) == 1.0


def _message_response(inner: dict[str, object], message_id: str) -> dict[str, object]:
    return {
        "id": inner["id"],
        "status": 200,
        "body": {
            "id": message_id,
            "subject": "",
            "bodyPreview": "",
            "isDraft": False,
            "body": {"contentType": "text", "content": message_id},
        },
    }


def _inner_message_id(inner: dict[str, object]) -> str:
    return str(inner["url"]).split("/messages/", 1)[1].split("?", 1)[0]


@pytest.mark.anyio
async def test_get_messages_retry_waits_share_one_budget(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    sleeps: list[float] = []

    async def fake_sleep(seconds: float) -> None:
        sleeps.append(seconds)

    monkeypatch.setattr(graph_module.anyio, "sleep", fake_sleep)

    def handler(request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        return httpx.Response(
            200,
            json={
                "responses": [
                    {"id": inner["id"], "status": 429, "headers": {"Retry-After": "8"}}
                    for inner in payload["requests"]
                ]
            },
        )

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    result = await graph.get_messages(messageIds=[f"m{index}" for index in range(50)])

    # Three chunks, each always throttled: the waits across all of them stay
    # inside one per-call budget instead of 3 x 3 x 30 s.
    # 8 + 8 fit the 20 s budget; a third 8 s wait would not, in this chunk
    # or the later ones, so they give up without sleeping.
    assert sleeps == [8.0, 8.0]
    assert all(item.error and "429" in item.error for item in result.messages)
    await client.aclose()


@pytest.mark.anyio
async def test_get_messages_keeps_other_chunks_when_a_batch_fails(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    sleeps: list[float] = []

    async def fake_sleep(seconds: float) -> None:
        sleeps.append(seconds)

    monkeypatch.setattr(graph_module.anyio, "sleep", fake_sleep)
    calls = {"second_chunk": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        ids = [_inner_message_id(inner) for inner in payload["requests"]]
        if "m25" in ids:
            calls["second_chunk"] += 1
            if calls["second_chunk"] == 1:
                # The whole $batch POST is throttled once, then succeeds.
                return httpx.Response(
                    429,
                    headers={"Retry-After": "2"},
                    json={"error": {"code": "TooManyRequests", "message": "slow down"}},
                )
        if "m45" in ids:
            return httpx.Response(
                500, json={"error": {"code": "InternalServerError", "message": "boom"}}
            )
        return httpx.Response(
            200,
            json={
                "responses": [
                    _message_response(inner, _inner_message_id(inner))
                    for inner in payload["requests"]
                ]
            },
        )

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    ids = [f"m{index}" for index in range(50)]
    result = await graph.get_messages(messageIds=ids)

    assert [item.messageId for item in result.messages] == ids
    # First chunk read normally; second chunk retried after an outer 429.
    assert all(item.message is not None for item in result.messages[:40])
    assert sleeps == [2.0]
    # Third chunk failed outright: per-item errors, nothing raised.
    for item in result.messages[40:]:
        assert item.message is None
        assert "500" in item.error and "InternalServerError" in item.error
    await client.aclose()


@pytest.mark.anyio
async def test_get_messages_rejects_empty_ids_and_survives_bad_items() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        responses = []
        for inner in payload["requests"]:
            message_id = _inner_message_id(inner)
            if message_id == "odd":
                # A 2xx that is not a message (e.g. a collection).
                responses.append({"id": inner["id"], "status": 200, "body": {"value": []}})
            else:
                responses.append(_message_response(inner, message_id))
        return httpx.Response(200, json={"responses": responses})

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    for bad in ([""], ["m1", "  "]):
        with pytest.raises(ValueError, match="empty"):
            await graph.get_messages(messageIds=bad)

    result = await graph.get_messages(messageIds=["m1", "odd", "m2"])
    assert result.messages[0].message.id == "m1"
    assert result.messages[1].message is None
    assert "Could not read message" in result.messages[1].error
    assert result.messages[2].message.id == "m2"
    await client.aclose()


@pytest.mark.anyio
async def test_get_message_returns_full_body_for_forwards() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        message_id = request.url.path.rsplit("/", 1)[1]
        subjects = {
            "fw": "FW: Rent roll",
            "fwd": "Fwd: Rent roll",
            "re-fw": "RE: [EXT] Fw: Rent roll",
            "vb": "VB: Hyreslista",
            "vs": "VS: Vuokralista",
            "wg": "AW: WG: Mietliste",
            "tr": "TR: Liste des loyers",
            "rv": "RV: Lista de alquileres",
            "reply": "RE: Rent roll forward plan",
        }
        return httpx.Response(
            200,
            json={
                "id": message_id,
                "subject": subjects[message_id],
                "bodyPreview": "",
                "isDraft": False,
                "body": {
                    "contentType": "text",
                    "content": "See below.\n---------- Forwarded message ----------\nRent roll attached.",
                },
                "uniqueBody": {"contentType": "text", "content": "See below."},
            },
        )

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    for message_id in ("fw", "fwd", "re-fw", "wg", "tr", "rv", "vb", "vs"):
        forwarded = await graph.get_message(messageId=message_id)
        assert forwarded.message.bodyScope == "full", message_id
        assert "Rent roll attached." in forwarded.message.body.content

    reply = await graph.get_message(messageId="reply")
    assert reply.message.bodyScope == "new"
    assert reply.message.body.content == "See below."

    await client.aclose()


@pytest.mark.anyio
async def test_request_refuses_absolute_urls_outside_graph() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        raise AssertionError(f"Unexpected request: {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    for url in (
        "https://evil.example/v1.0/me/messages?$skiptoken=abc",
        "https://graph.microsoft.com.evil.example/v1.0/me/messages",
    ):
        with pytest.raises(ValueError, match="graph.microsoft.com"):
            await graph._request(url)

    await client.aclose()


@pytest.mark.anyio
async def test_get_thread_sorts_and_limits_on_graph_in_one_request() -> None:
    def message(index: int) -> dict[str, object]:
        return {
            "id": f"t{index}",
            "subject": "Deal",
            "bodyPreview": "",
            "isDraft": False,
            "conversationId": "conv-9",
            "receivedDateTime": f"2026-05-{index + 1:02d}T12:00:00Z",
        }

    thread = [message(index) for index in range(5)]
    requests: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        params = request.url.params
        # The sorted property must lead the filter for Graph to accept $orderby.
        assert params["$filter"] == (
            "receivedDateTime ge 1900-01-01T00:00:00Z and conversationId eq 'conv-9'"
        )
        top = int(params["$top"])
        ordered = sorted(
            thread,
            key=lambda item: str(item["receivedDateTime"]),
            reverse=params["$orderby"] == "receivedDateTime desc",
        )
        body: dict[str, object] = {"value": ordered[:top]}
        if top < len(ordered):
            body["@odata.nextLink"] = (
                "https://graph.microsoft.com/v1.0/me/messages?$skiptoken=abc"
            )
        return httpx.Response(200, json=body)

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    latest = await graph.get_thread(conversationId="conv-9", top=2)
    assert len(requests) == 1
    assert requests[-1].url.params["$top"] == "2"
    assert requests[-1].url.params["$orderby"] == "receivedDateTime desc"
    assert [item.id for item in latest.messages] == ["t4", "t3"]
    assert latest.order == "newest"
    # More messages exist, so the size is unknown rather than reported as 2.
    assert latest.truncated is True
    assert latest.messageCount is None

    earliest = await graph.get_thread(conversationId="conv-9", top=10, order="oldest")
    assert len(requests) == 2
    assert requests[-1].url.params["$orderby"] == "receivedDateTime asc"
    assert [item.id for item in earliest.messages] == ["t0", "t1", "t2", "t3", "t4"]
    assert earliest.truncated is False
    assert earliest.messageCount == 5

    await graph.get_thread(conversationId="conv-9", top=500)
    assert requests[-1].url.params["$top"] == "100"

    with pytest.raises(ValueError, match="order"):
        await graph.get_thread(conversationId="conv-9", order="latest")

    await client.aclose()


@pytest.mark.anyio
async def test_create_send_and_move_message() -> None:
    requests: list[tuple[str, str, dict[str, object] | None]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content.decode("utf-8")) if request.content else None
        requests.append((request.method, request.url.path, body))

        if request.method == "POST" and request.url.path.endswith("/messages"):
            return httpx.Response(
                200,
                json={
                    "id": "draft-99",
                    "subject": "Created",
                    "from": {"emailAddress": {"address": "delegated@example.com"}},
                    "bodyPreview": "Draft preview",
                    "isDraft": True,
                },
            )

        if request.method == "POST" and request.url.path.endswith("/messages/draft-99/send"):
            return httpx.Response(204)

        if request.url.path.endswith("/mailFolders('Archive')"):
            return httpx.Response(200, json={"id": "folder-archive", "displayName": "Archive"})

        if request.method == "POST" and request.url.path.endswith("/messages/mail-1/move"):
            return httpx.Response(
                200,
                json={"id": "mail-1", "subject": "Moved", "bodyPreview": "Moved preview"},
            )

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    draft = await graph.create_draft(
        mailbox="shared@example.com",
        subject="Created",
        to=["a@example.com"],
        cc=["b@example.com"],
        bcc=["c@example.com"],
        body="Hello",
        bodyType="html",
        from_="delegated@example.com",
    )
    assert draft.draft.id == "draft-99"

    sent = await graph.send_draft(mailbox="shared@example.com", messageId="draft-99")
    assert sent.sent is True

    moved = await graph.move_message(
        mailbox="shared@example.com",
        messageId="mail-1",
        destinationFolder="Archive",
    )
    assert moved.destinationFolder == "Archive"

    create_body = requests[0][2]
    assert create_body is not None
    assert create_body["from"] == {"emailAddress": {"address": "delegated@example.com"}}
    assert create_body["body"] == {"contentType": "HTML", "content": "Hello"}
    assert create_body["toRecipients"] == [
        {"emailAddress": {"address": "a@example.com"}}
    ]
    assert create_body["ccRecipients"] == [
        {"emailAddress": {"address": "b@example.com"}}
    ]
    assert create_body["bccRecipients"] == [
        {"emailAddress": {"address": "c@example.com"}}
    ]

    move_body = requests[-1][2]
    assert move_body == {"destinationId": "folder-archive"}

    await client.aclose()


@pytest.mark.anyio
async def test_direct_send_and_reply() -> None:
    requests: list[tuple[str, str, dict[str, object] | None]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content.decode("utf-8")) if request.content else None
        requests.append((request.method, request.url.path, body))

        if request.method == "POST" and request.url.path.endswith("/sendMail"):
            assert body == {
                "message": {
                    "subject": "Quick note",
                    "body": {"contentType": "Text", "content": "Hello"},
                    "toRecipients": [
                        {"emailAddress": {"address": "a@example.com"}}
                    ],
                },
                "saveToSentItems": True,
            }
            return httpx.Response(202)

        if request.method == "POST" and request.url.path.endswith("/messages/msg-1/replyAll"):
            assert body == {
                "message": {
                    "body": {"contentType": "HTML", "content": "<p>Thanks</p>"}
                }
            }
            return httpx.Response(202)

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    sent = await graph.send_mail(
        subject="Quick note",
        to=["a@example.com"],
        body="Hello",
    )
    assert sent.sent is True
    assert sent.subject == "Quick note"

    replied = await graph.send_reply(
        messageId="msg-1",
        comment="<p>Thanks</p>",
        replyAll=True,
    )
    assert replied.sent is True
    assert replied.messageId == "msg-1"
    assert len(requests) == 2

    await client.aclose()


@pytest.mark.anyio
async def test_create_draft_omits_empty_optional_fields() -> None:
    captured_body: dict[str, object] | None = None

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal captured_body

        if request.method == "POST" and request.url.path.endswith("/messages"):
            captured_body = json.loads(request.content.decode("utf-8"))
            return httpx.Response(
                200,
                json={
                    "id": "draft-minimal",
                    "subject": "Minimal",
                    "bodyPreview": "Draft preview",
                    "isDraft": True,
                },
            )

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    draft = await graph.create_draft(
        subject="Minimal",
        to=[],
        cc=None,
        bcc=None,
        body="Hello",
        bodyType="text",
        from_=None,
    )

    assert draft.draft.id == "draft-minimal"
    assert captured_body == {
        "subject": "Minimal",
        "body": {"contentType": "Text", "content": "Hello"},
    }

    await client.aclose()


@pytest.mark.anyio
async def test_folder_navigation_inbox_filters_and_nested_move() -> None:
    requests: list[tuple[str, str, dict[str, str], dict[str, object] | None]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content.decode("utf-8")) if request.content else None
        requests.append(
            (request.method, request.url.path, dict(request.url.params), body)
        )

        if request.url.path.endswith("/mailFolders") and request.method == "GET":
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "inbox-id",
                            "displayName": "Inbox",
                            "childFolderCount": 1,
                            "totalItemCount": 10,
                            "unreadItemCount": 2,
                        }
                    ]
                },
            )

        if request.url.path.endswith("/mailFolders/inbox-id/childFolders"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "clients-id",
                            "displayName": "Clients",
                            "parentFolderId": "inbox-id",
                            "childFolderCount": 1,
                        }
                    ]
                },
            )

        if request.url.path.endswith("/mailFolders/clients-id/childFolders"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "acme-id",
                            "displayName": "Acme",
                            "parentFolderId": "clients-id",
                            "childFolderCount": 0,
                        }
                    ]
                },
            )

        if request.url.path.endswith("/mailFolders/acme-id/messages") and request.method == "GET":
            assert "inferenceClassification" in request.url.params["$select"]
            assert request.url.params["$orderby"] == "receivedDateTime desc"
            assert request.url.params["$filter"] == (
                "receivedDateTime ge 1900-01-01T00:00:00Z and "
                "isRead eq false and hasAttachments eq true and "
                "importance eq 'high' and categories/any(c:c eq 'Client') and "
                "flag/flagStatus eq 'flagged' and "
                "inferenceClassification eq 'focused'"
            )
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "m-nav",
                            "subject": "Needs review",
                            "from": {"emailAddress": {"address": "client@example.com"}},
                            "sender": {"emailAddress": {"address": "assistant@example.com"}},
                            "replyTo": [{"emailAddress": {"address": "reply@example.com"}}],
                            "bodyPreview": "Please review",
                            "isDraft": False,
                            "isRead": False,
                            "hasAttachments": True,
                            "importance": "high",
                            "categories": ["Client"],
                            "flag": {"flagStatus": "flagged"},
                            "inferenceClassification": "focused",
                            "parentFolderId": "acme-id",
                            "internetMessageId": "<message@example.com>",
                            "conversationId": "conv-nav",
                        }
                    ]
                },
            )

        if request.url.path.endswith("/messages/m-nav/move"):
            return httpx.Response(
                200,
                json={"id": "m-nav", "subject": "Moved", "bodyPreview": "", "isDraft": False},
            )

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    tree = await graph.mail_folder_tree(maxDepth=3)
    assert tree.folders[0].childFolders[0].childFolders[0].path == "Inbox/Clients/Acme"

    resolved = await graph.resolve_mail_folder(folderPath="Inbox/Clients/Acme")
    assert resolved.folder.id == "acme-id"

    listed = await graph.list_messages(
        folderPath="Inbox/Clients/Acme",
        isRead=False,
        hasAttachments=True,
        importance="high",
        categories=["Client"],
        flagStatus="flagged",
        inferenceClassification="focused",
    )
    message = listed.messages[0]
    assert message.sender == "assistant@example.com"
    assert message.replyTo == ["reply@example.com"]
    assert message.isRead is False
    assert message.hasAttachments is True
    assert message.categories == ["Client"]
    assert message.flagStatus == "flagged"
    assert message.inferenceClassification == "focused"
    assert message.parentFolderId == "acme-id"

    moved = await graph.move_message(
        messageId="m-nav",
        destinationFolder="Inbox/Clients/Acme",
    )
    assert moved.destinationFolderId == "acme-id"
    assert requests[-1][3] == {"destinationId": "acme-id"}

    await client.aclose()


@pytest.mark.anyio
async def test_folder_mutations_and_rules() -> None:
    requests: list[tuple[str, str, dict[str, object] | None]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content.decode("utf-8")) if request.content else None
        requests.append((request.method, request.url.path, body))

        if request.url.path.endswith("/mailFolders") and request.method == "GET":
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "inbox-id",
                            "displayName": "Inbox",
                            "childFolderCount": 1,
                        }
                    ]
                },
            )

        if request.url.path.endswith("/mailFolders/inbox-id/childFolders") and request.method == "GET":
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "clients-id",
                            "displayName": "Clients",
                            "parentFolderId": "inbox-id",
                            "childFolderCount": 0,
                        }
                    ]
                },
            )

        if request.url.path.endswith("/mailFolders/clients-id/childFolders") and request.method == "POST":
            assert body == {"displayName": "Acme"}
            return httpx.Response(
                201,
                json={
                    "id": "acme-id",
                    "displayName": "Acme",
                    "parentFolderId": "clients-id",
                    "childFolderCount": 0,
                },
            )

        if request.url.path.endswith("/mailFolders/acme-id") and request.method == "PATCH":
            assert body == {"displayName": "Acme Corp"}
            return httpx.Response(
                200,
                json={
                    "id": "acme-id",
                    "displayName": "Acme Corp",
                    "parentFolderId": "clients-id",
                    "childFolderCount": 0,
                },
            )

        if request.url.path.endswith("/mailFolders/acme-id") and request.method == "DELETE":
            return httpx.Response(204)

        if request.url.path.endswith("/mailFolders/inbox/messageRules") and request.method == "GET":
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "rule-1",
                            "displayName": "Clients",
                            "sequence": 1,
                            "isEnabled": True,
                            "actions": {"markAsRead": True},
                            "conditions": {"senderContains": ["client"]},
                        }
                    ]
                },
            )

        if request.url.path.endswith("/mailFolders/inbox/messageRules") and request.method == "POST":
            assert body == {
                "displayName": "Move Acme",
                "sequence": 2,
                "isEnabled": True,
                "conditions": {
                    "senderContains": ["acme"],
                    "subjectContains": ["Invoice"],
                },
                "actions": {
                    "moveToFolder": "clients-id",
                    "markAsRead": True,
                    "assignCategories": ["Client"],
                },
            }
            return httpx.Response(
                201,
                json={
                    "id": "rule-2",
                    "displayName": "Move Acme",
                    "sequence": 2,
                    "isEnabled": True,
                    "actions": body["actions"],
                    "conditions": body["conditions"],
                },
            )

        if request.url.path.endswith("/mailFolders/inbox/messageRules/rule-2") and request.method == "PATCH":
            assert body == {"displayName": "Move Acme invoices", "isEnabled": False}
            return httpx.Response(
                200,
                json={
                    "id": "rule-2",
                    "displayName": "Move Acme invoices",
                    "sequence": 2,
                    "isEnabled": False,
                },
            )

        if request.url.path.endswith("/mailFolders/inbox/messageRules/rule-2") and request.method == "DELETE":
            return httpx.Response(204)

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    created_folder = await graph.create_mail_folder(
        displayName="Acme",
        parentFolderPath="Inbox/Clients",
    )
    assert created_folder.folder.id == "acme-id"

    renamed_folder = await graph.rename_mail_folder(
        folderId="acme-id",
        displayName="Acme Corp",
    )
    assert renamed_folder.folder.displayName == "Acme Corp"

    deleted_folder = await graph.delete_mail_folder(folderId="acme-id")
    assert deleted_folder.deleted is True

    rules = await graph.list_mail_rules()
    assert rules.rules[0].displayName == "Clients"

    created_rule = await graph.create_mail_rule(
        displayName="Move Acme",
        sequence=2,
        senderContains=["acme"],
        subjectContains=["Invoice"],
        moveToFolderPath="Inbox/Clients",
        markAsRead=True,
        assignCategories=["Client"],
    )
    assert created_rule.rule.id == "rule-2"

    updated_rule = await graph.update_mail_rule(
        ruleId="rule-2",
        displayName="Move Acme invoices",
        isEnabled=False,
    )
    assert updated_rule.rule.isEnabled is False

    deleted_rule = await graph.delete_mail_rule(ruleId="rule-2")
    assert deleted_rule.deleted is True

    await client.aclose()


@pytest.mark.anyio
async def test_attachments_threads_and_categories() -> None:
    text_payload = base64.b64encode(b"hello attachment").decode("ascii")
    requests: list[tuple[str, str, dict[str, str], dict[str, object] | None]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content.decode("utf-8")) if request.content else None
        requests.append((request.method, request.url.path, dict(request.url.params), body))

        if request.url.path.endswith("/messages/msg-1/attachments") and request.method == "GET":
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "@odata.type": "#microsoft.graph.fileAttachment",
                            "id": "att-1",
                            "name": "notes.txt",
                            "contentType": "text/plain",
                            "size": 16,
                            "isInline": False,
                        },
                        {
                            "@odata.type": "#microsoft.graph.fileAttachment",
                            "id": "att-inline",
                            "name": "pixel.png",
                            "contentType": "image/png",
                            "size": 10,
                            "isInline": True,
                        },
                    ]
                },
            )

        if request.url.path.endswith("/messages/msg-1/attachments/att-1/$value"):
            return httpx.Response(200, content=base64.b64decode(text_payload))

        if request.url.path.endswith("/messages/msg-1/attachments/att-1"):
            # Metadata only: contentBytes would download the file before the
            # size check.
            assert "contentBytes" not in request.url.params["$select"]
            return httpx.Response(
                200,
                json={
                    "@odata.type": "#microsoft.graph.fileAttachment",
                    "id": "att-1",
                    "name": "notes.txt",
                    "contentType": "text/plain",
                    "size": 16,
                },
            )

        if request.url.path.endswith("/messages/msg-1/attachments/bin-1"):
            return httpx.Response(
                200,
                json={
                    "@odata.type": "#microsoft.graph.fileAttachment",
                    "id": "bin-1",
                    "name": "image.png",
                    "contentType": "image/png",
                    "size": 16,
                },
            )

        if request.url.path.endswith("/messages/msg-1/createReplyAll"):
            assert body == {
                "message": {
                    "body": {"contentType": "HTML", "content": "<p>Thanks</p>"}
                }
            }
            return httpx.Response(
                201,
                json={
                    "id": "reply-draft",
                    "subject": "Re: Hello",
                    "bodyPreview": "Thanks",
                    "isDraft": True,
                    "conversationId": "conv-1",
                },
            )

        if request.url.path.endswith("/messages/msg-1") and request.method == "GET":
            assert request.url.params["$select"] == "conversationId"
            return httpx.Response(200, json={"conversationId": "conv-1"})

        if request.url.path.endswith(
            "/messages"
        ) and "conversationId" in request.url.params.get("$filter", ""):
            assert request.url.params["$orderby"].startswith("receivedDateTime ")
            thread = [
                {
                    "id": "msg-2",
                    "subject": "Follow-up",
                    "bodyPreview": "Second",
                    "isDraft": False,
                    "conversationId": "conv-1",
                    "receivedDateTime": "2026-05-13T12:05:00Z",
                },
                {
                    "id": "msg-1",
                    "subject": "Hello",
                    "bodyPreview": "First",
                    "isDraft": False,
                    "conversationId": "conv-1",
                    "receivedDateTime": "2026-05-13T12:00:00Z",
                },
            ]
            if request.url.params["$orderby"].endswith(" asc"):
                thread.reverse()
            return httpx.Response(
                200,
                json={"value": thread},
            )

        if request.url.path.endswith("/outlook/masterCategories") and request.method == "GET":
            return httpx.Response(
                200,
                json={"value": [{"id": "cat-1", "displayName": "Client", "color": "preset1"}]},
            )

        if request.url.path.endswith("/outlook/masterCategories") and request.method == "POST":
            return httpx.Response(
                201,
                json={"id": "cat-2", "displayName": body["displayName"], "color": body["color"]},
            )

        if request.url.path.endswith("/outlook/masterCategories/cat-2") and request.method == "PATCH":
            assert body == {"color": "preset3"}
            return httpx.Response(
                200,
                json={"id": "cat-2", "displayName": "Prospect", "color": body["color"]},
            )

        if request.url.path.endswith("/outlook/masterCategories/cat-2") and request.method == "DELETE":
            return httpx.Response(204)

        if request.url.path.endswith("/messages/msg-1") and request.method == "PATCH":
            return httpx.Response(
                200,
                json={
                    "id": "msg-1",
                    "subject": "Hello",
                    "bodyPreview": "",
                    "isDraft": False,
                    "categories": body.get("categories", []),
                    "flag": body.get("flag", {}),
                    "isRead": body.get("isRead"),
                },
            )

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    attachments = await graph.list_attachments(messageId="msg-1")
    assert [attachment.id for attachment in attachments.attachments] == ["att-1"]

    content = await graph.get_attachment_content(messageId="msg-1", attachmentId="att-1")
    assert content.content == "hello attachment"

    binary = await graph.get_attachment_content(messageId="msg-1", attachmentId="bin-1")
    assert binary.content is None
    assert binary.unsupportedReason is not None

    thread = await graph.get_thread(conversationId="conv-1")
    assert thread.order == "newest"
    assert [message.id for message in thread.messages] == ["msg-2", "msg-1"]

    thread_by_message = await graph.get_thread(messageId="msg-1", order="oldest")
    assert thread_by_message.conversationId == "conv-1"
    assert [message.id for message in thread_by_message.messages] == ["msg-1", "msg-2"]

    reply = await graph.create_reply_draft(
        messageId="msg-1",
        comment="<p>Thanks</p>",
        replyAll=True,
    )
    assert reply.draft.id == "reply-draft"

    categories = await graph.list_categories()
    assert categories.categories[0].displayName == "Client"

    created = await graph.create_category(displayName="Prospect", color="preset2")
    assert created.category.displayName == "Prospect"

    updated = await graph.update_category(categoryId="cat-2", color="preset3")
    assert updated.category.color == "preset3"

    with pytest.raises(ValueError, match="does not support renaming"):
        await graph.update_category(categoryId="cat-2", displayName="Customer")

    deleted = await graph.delete_category(categoryId="cat-2")
    assert deleted.deleted is True

    categorized = await graph.set_message_categories(
        messageId="msg-1",
        categories=["Client"],
    )
    assert categorized.message.categories == ["Client"]

    read = await graph.mark_message_read(messageId="msg-1", isRead=True)
    assert read.message.isRead is True

    flagged = await graph.set_message_flag(messageId="msg-1", flagStatus="flagged")
    assert flagged.message.flagStatus == "flagged"

    await client.aclose()


@pytest.mark.anyio
async def test_attachment_images_and_inline_pictures() -> None:
    png_payload = base64.b64encode(b"\x89PNG\r\n\x1a\nlogo").decode("ascii")
    jpeg_bytes = b"\xff\xd8\xff\xe0jpeg-body"
    requested_paths: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested_paths.append(request.url.path)

        if request.url.path.endswith("/messages/msg-1/attachments"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "@odata.type": "#microsoft.graph.fileAttachment",
                            "id": "inline-png",
                            "name": "logo.png",
                            "contentType": "image/png",
                            "size": 12,
                            "isInline": True,
                            "contentId": "logo@01D9",
                            "contentBytes": png_payload,
                        },
                        {
                            # Inline by content ID only, and Graph omitted
                            # contentBytes so the bytes come from /$value.
                            "@odata.type": "#microsoft.graph.fileAttachment",
                            "id": "inline-jpg",
                            "name": "signature.JPG",
                            "contentType": "application/octet-stream",
                            "size": len(jpeg_bytes),
                            "isInline": False,
                            "contentId": "sig@01D9",
                        },
                        {
                            "@odata.type": "#microsoft.graph.fileAttachment",
                            "id": "inline-huge",
                            "name": "banner.png",
                            "contentType": "image/png",
                            "size": 9_000_000,
                            "isInline": True,
                            "contentId": "banner@01D9",
                        },
                        {
                            "@odata.type": "#microsoft.graph.fileAttachment",
                            "id": "attached-png",
                            "name": "chart.png",
                            "contentType": "image/png",
                            "size": 12,
                            "isInline": False,
                            "contentBytes": png_payload,
                        },
                        {
                            "@odata.type": "#microsoft.graph.fileAttachment",
                            "id": "notes",
                            "name": "notes.txt",
                            "contentType": "text/plain",
                            "size": 4,
                            "isInline": False,
                        },
                    ]
                },
            )

        if request.url.path.endswith("/messages/msg-1/attachments/inline-jpg/$value"):
            return httpx.Response(200, content=jpeg_bytes)

        if request.url.path.endswith("/messages/msg-1/attachments/inline-png/$value"):
            return httpx.Response(200, content=base64.b64decode(png_payload))

        if request.url.path.endswith("/messages/msg-1/attachments/inline-png"):
            assert "contentBytes" not in request.url.params["$select"]
            return httpx.Response(
                200,
                json={
                    "@odata.type": "#microsoft.graph.fileAttachment",
                    "id": "inline-png",
                    "name": "logo.png",
                    "contentType": "image/png",
                    "size": 12,
                    "isInline": True,
                    "contentId": "logo@01D9",
                },
            )

        if request.url.path.endswith("/messages/msg-1/attachments/notes"):
            return httpx.Response(
                200,
                json={
                    "@odata.type": "#microsoft.graph.fileAttachment",
                    "id": "notes",
                    "name": "notes.txt",
                    "contentType": "text/plain",
                    "size": 4,
                },
            )

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    single = await graph.get_attachment_image(
        messageId="msg-1",
        attachmentId="inline-png",
    )
    assert single.unsupportedReason is None
    assert single.image is not None
    assert single.image.mimeType == "image/png"
    assert single.image.dataBase64 == png_payload
    assert single.image.byteSize == 12
    assert single.image.attachment.contentId == "logo@01D9"

    not_an_image = await graph.get_attachment_image(
        messageId="msg-1",
        attachmentId="notes",
    )
    assert not_an_image.image is None
    assert "not a PNG" in not_an_image.unsupportedReason

    inline = await graph.get_inline_images(messageId="msg-1")
    assert [image.attachment.id for image in inline.images] == [
        "inline-png",
        "inline-jpg",
    ]
    # Content type was generic, so the mime type comes from the file extension.
    assert inline.images[1].mimeType == "image/jpeg"
    assert base64.b64decode(inline.images[1].dataBase64) == jpeg_bytes
    assert [skipped.attachment.id for skipped in inline.skipped] == ["inline-huge"]
    assert "maxBytes" in inline.skipped[0].reason
    assert inline.truncated is False

    with_attachments = await graph.get_inline_images(
        messageId="msg-1",
        includeNonInline=True,
    )
    assert [image.attachment.id for image in with_attachments.images] == [
        "inline-png",
        "inline-jpg",
        "attached-png",
    ]

    capped = await graph.get_inline_images(messageId="msg-1", maxImages=1)
    assert [image.attachment.id for image in capped.images] == ["inline-png"]
    assert capped.truncated is True

    # The batch budget stops before a picture that would overflow it, even
    # though each picture on its own is under maxBytes.
    budgeted = await graph.get_inline_images(messageId="msg-1", maxTotalBytes=14)
    assert [image.attachment.id for image in budgeted.images] == ["inline-png"]
    assert budgeted.truncated is True

    # Text attachment reads point at the image tools instead of failing silently.
    image_via_text_tool = await graph.get_attachment_content(
        messageId="msg-1",
        attachmentId="inline-png",
    )
    assert image_via_text_tool.content is None
    assert "mail_get_attachment_image" in image_via_text_tool.unsupportedReason

    assert "/messages/msg-1/attachments/inline-huge/$value" not in requested_paths

    await client.aclose()


@pytest.mark.anyio
async def test_pdf_attachment_text_extraction(monkeypatch: pytest.MonkeyPatch) -> None:
    class FakePage:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class FakePdfReader:
        def __init__(self, stream: object) -> None:
            self.pages = [FakePage("First page"), FakePage("Second page")]

    monkeypatch.setattr(graph_module, "PdfReader", FakePdfReader)
    pdf_payload = base64.b64encode(b"%PDF fake content").decode("ascii")

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/messages/msg-1/attachments/pdf-1/$value"):
            return httpx.Response(200, content=base64.b64decode(pdf_payload))

        if request.url.path.endswith("/messages/msg-1/attachments/pdf-1"):
            assert "contentBytes" not in request.url.params["$select"]
            return httpx.Response(
                200,
                json={
                    "@odata.type": "#microsoft.graph.fileAttachment",
                    "id": "pdf-1",
                    "name": "brief.pdf",
                    "contentType": "application/pdf",
                    "size": 128,
                },
            )

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    content = await graph.get_attachment_content(
        messageId="msg-1",
        attachmentId="pdf-1",
        maxChars=20,
    )

    assert content.encoding == "pdf-text"
    assert content.content == "--- Page 1 ---\nFirst"
    assert content.truncated is True
    assert "maxChars=20" in content.unsupportedReason

    await client.aclose()


def _single_page_pdf(text: bytes = b"INVOICE 12345", pages: int = 2) -> bytes:
    """Build a small valid PDF so rendering runs against real pdfium."""

    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # placeholder for the page tree, filled in below
    ]
    kids: list[bytes] = []
    for page_index in range(pages):
        page_obj = len(objects) + 1
        content_obj = page_obj + 1
        kids.append(b"%d 0 R" % page_obj)
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
            % (2 + pages * 2 + 1, content_obj)
        )
        stream = b"BT /F1 36 Tf 72 700 Td (%s p%d) Tj ET" % (text, page_index + 1)
        objects.append(
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        )
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(kids), pages)
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


@pytest.mark.anyio
async def test_pdf_attachment_page_rendering() -> None:
    pdf_bytes = _single_page_pdf(pages=3)

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/messages/msg-1/attachments/pdf-1/$value"):
            return httpx.Response(200, content=pdf_bytes)

        if request.url.path.endswith("/messages/msg-1/attachments/pdf-1"):
            assert "contentBytes" not in request.url.params["$select"]
            return httpx.Response(
                200,
                json={
                    "@odata.type": "#microsoft.graph.fileAttachment",
                    "id": "pdf-1",
                    "name": "invoice.pdf",
                    "contentType": "application/pdf",
                    "size": len(pdf_bytes),
                },
            )

        if request.url.path.endswith("/messages/msg-1/attachments/notes"):
            return httpx.Response(
                200,
                json={
                    "@odata.type": "#microsoft.graph.fileAttachment",
                    "id": "notes",
                    "name": "notes.txt",
                    "contentType": "text/plain",
                    "size": 4,
                },
            )

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    rendered = await graph.get_attachment_pdf_pages(
        messageId="msg-1",
        attachmentId="pdf-1",
        maxPages=2,
    )
    assert rendered.unsupportedReason is None
    assert rendered.pageCount == 3
    assert [page.pageNumber for page in rendered.pages] == [1, 2]
    assert rendered.truncated is True
    first = rendered.pages[0]
    assert first.mimeType == "image/jpeg"
    assert first.byteSize > 0
    # Long edge honours the requested target, aspect ratio preserved.
    assert max(first.widthPx, first.heightPx) == 1600
    assert base64.b64decode(first.dataBase64)[:2] == b"\xff\xd8"

    paged = await graph.get_attachment_pdf_pages(
        messageId="msg-1",
        attachmentId="pdf-1",
        firstPage=3,
    )
    assert [page.pageNumber for page in paged.pages] == [3]
    assert paged.truncated is False

    smaller = await graph.get_attachment_pdf_pages(
        messageId="msg-1",
        attachmentId="pdf-1",
        maxPages=1,
        longEdge=400,
    )
    assert max(smaller.pages[0].widthPx, smaller.pages[0].heightPx) == 400
    assert smaller.pages[0].byteSize < first.byteSize

    # The byte budget still yields the first page rather than nothing.
    budgeted = await graph.get_attachment_pdf_pages(
        messageId="msg-1",
        attachmentId="pdf-1",
        maxTotalBytes=1,
    )
    assert [page.pageNumber for page in budgeted.pages] == [1]
    assert budgeted.truncated is True

    beyond_end = await graph.get_attachment_pdf_pages(
        messageId="msg-1",
        attachmentId="pdf-1",
        firstPage=99,
    )
    assert beyond_end.pages == []
    assert "has 3 pages" in beyond_end.unsupportedReason

    not_pdf = await graph.get_attachment_pdf_pages(
        messageId="msg-1",
        attachmentId="notes",
    )
    assert not_pdf.pages == []
    assert not_pdf.unsupportedReason == "Attachment is not a PDF"

    await client.aclose()


@pytest.mark.anyio
async def test_pdf_page_rendering_reports_damaged_file() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/$value"):
            return httpx.Response(200, content=b"not a pdf")
        return httpx.Response(
            200,
            json={
                "@odata.type": "#microsoft.graph.fileAttachment",
                "id": "pdf-bad",
                "name": "broken.pdf",
                "contentType": "application/pdf",
                "size": 9,
            },
        )

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    result = await graph.get_attachment_pdf_pages(
        messageId="msg-1",
        attachmentId="pdf-bad",
    )
    assert result.pages == []
    assert "Could not open PDF" in result.unsupportedReason

    await client.aclose()


@pytest.mark.anyio
async def test_pdf_page_rendering_without_renderer(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(graph_module, "pypdfium2", None)

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "@odata.type": "#microsoft.graph.fileAttachment",
                "id": "pdf-1",
                "name": "invoice.pdf",
                "contentType": "application/pdf",
                "size": 128,
            },
        )

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    result = await graph.get_attachment_pdf_pages(
        messageId="msg-1",
        attachmentId="pdf-1",
    )
    assert result.pages == []
    assert "pypdfium2" in result.unsupportedReason

    await client.aclose()


@pytest.mark.anyio
async def test_contacts_crud_search_and_folders() -> None:
    requests: list[tuple[str, str, dict[str, str], dict[str, object] | None]] = []
    contact_categories = ["Client"]
    contact_display_name = ["Ada Lovelace"]

    def contact(contact_id: str = "contact-1", name: str = "Ada Lovelace") -> dict[str, object]:
        return {
            "id": contact_id,
            "displayName": name,
            "givenName": "Ada",
            "surname": "Lovelace",
            "companyName": "Analytical Engines",
            "jobTitle": "Mathematician",
            "personalNotes": "Analytical notes for client follow-up",
            "singleValueExtendedProperties": [
                {"id": "String 0x3A50", "value": "https://analytical.example"}
            ],
            "businessPhones": ["555-0100"],
            "mobilePhone": "555-0101",
            "emailAddresses": [{"address": "ada@example.com", "name": name}],
            "categories": contact_categories,
            "parentFolderId": "contacts-folder",
            "businessAddress": {
                "street": "1 Analytical Way",
                "city": "London",
                "state": "",
                "countryOrRegion": "UK",
                "postalCode": "NW1",
            },
            "homeAddress": {},
            "otherAddress": {
                "street": "PO Box 1",
                "city": "Seattle",
                "state": "WA",
                "countryOrRegion": "US",
                "postalCode": "98101",
            },
        }

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content.decode("utf-8")) if request.content else None
        requests.append((request.method, request.url.path, dict(request.url.params), body))

        if request.url.path.endswith("/contactFolders") and request.method == "GET":
            assert request.url.params["$select"] == "id,displayName,parentFolderId"
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "contacts-folder",
                            "displayName": "VIP",
                        }
                    ]
                },
            )

        if request.url.path.endswith("/contacts") and request.method == "GET":
            assert "categories" in request.url.params["$select"]
            assert "businessAddress" in request.url.params["$select"]
            assert "businessHomePage" not in request.url.params["$select"]
            assert "personalNotes" in request.url.params["$select"]
            assert request.url.params["$expand"] == (
                "singleValueExtendedProperties($filter=id eq 'String 0x3A50')"
            )
            if "$filter" in request.url.params:
                assert "emailAddresses/any" in request.url.params["$filter"]
            return httpx.Response(200, json={"value": [contact()]})

        if request.url.path.endswith("/contacts/contact-1") and request.method == "GET":
            assert request.url.params["$expand"] == (
                "singleValueExtendedProperties($filter=id eq 'String 0x3A50')"
            )
            return httpx.Response(
                200, json=contact("contact-1", contact_display_name[0])
            )

        if request.url.path.endswith("/contacts/contact-new") and request.method == "GET":
            assert request.url.params["$expand"] == (
                "singleValueExtendedProperties($filter=id eq 'String 0x3A50')"
            )
            return httpx.Response(200, json=contact("contact-new"))

        if request.url.path.endswith("/contacts") and request.method == "POST":
            assert body["emailAddresses"] == [
                {"address": "ada@example.com", "name": "Ada Lovelace"}
            ]
            assert body["categories"] == ["Client", "VIP"]
            assert body["personalNotes"] == "Met at the analytics conference"
            assert body["singleValueExtendedProperties"] == [
                {"id": "String 0x3A50", "value": "https://ada.example"}
            ]
            assert body["businessAddress"] == {
                "street": "1 Analytical Way",
                "city": "London",
                "countryOrRegion": "UK",
                "postalCode": "NW1",
            }
            return httpx.Response(201, json=contact("contact-new"))

        if request.url.path.endswith("/contacts/contact-1") and request.method == "PATCH":
            if "categories" in body:
                contact_categories[:] = body["categories"]
            if "displayName" in body:
                assert body["personalNotes"] == "Updated relationship notes"
                assert body["singleValueExtendedProperties"] == [
                    {"id": "String 0x3A50", "value": "https://byron.example"}
                ]
                contact_display_name[0] = body["displayName"]
            return httpx.Response(
                200,
                json=contact("contact-1", contact_display_name[0]),
            )

        if request.url.path.endswith("/contacts/contact-1") and request.method == "DELETE":
            return httpx.Response(204)

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    folders = await graph.list_contact_folders(mailbox="shared@example.com")
    assert folders.folders[0].displayName == "VIP"
    assert folders.folders[0].childFolderCount == 0

    listed = await graph.list_contacts(mailbox="shared@example.com")
    assert listed.contacts[0].emailAddresses == ["ada@example.com"]
    assert listed.contacts[0].categories == ["Client"]
    assert listed.contacts[0].parentFolderId == "contacts-folder"
    assert listed.contacts[0].personalHomePage == "https://analytical.example"
    assert listed.contacts[0].personalNotes == "Analytical notes for client follow-up"
    assert listed.contacts[0].businessAddress.city == "London"
    assert listed.contacts[0].homeAddress is None
    assert listed.contacts[0].otherAddress.state == "WA"

    searched = await graph.search_contacts(
        mailbox="shared@example.com",
        query="ada@example.com",
    )
    assert searched.contacts[0].displayName == "Ada Lovelace"

    searched_note = await graph.search_contacts(
        mailbox="shared@example.com",
        query="follow-up",
    )
    assert searched_note.contacts[0].personalNotes == "Analytical notes for client follow-up"

    searched_home_page = await graph.search_contacts(
        mailbox="shared@example.com",
        query="analytical.example",
    )
    assert searched_home_page.contacts[0].personalHomePage == "https://analytical.example"

    got = await graph.get_contact(
        mailbox="shared@example.com",
        contactId="contact-1",
    )
    assert got.contact.companyName == "Analytical Engines"
    assert got.contact.personalHomePage == "https://analytical.example"

    created = await graph.create_contact(
        mailbox="shared@example.com",
        displayName="Ada Lovelace",
        emailAddresses=["ada@example.com"],
        categories=["Client", "VIP"],
        personalHomePage="https://ada.example",
        personalNotes="Met at the analytics conference",
        businessAddress={
            "street": "1 Analytical Way",
            "city": "London",
            "state": None,
            "countryOrRegion": "UK",
            "postalCode": "NW1",
        },
    )
    assert created.contact.id == "contact-new"

    updated = await graph.update_contact(
        mailbox="shared@example.com",
        contactId="contact-1",
        displayName="Ada Byron",
        personalHomePage="https://byron.example",
        personalNotes="Updated relationship notes",
        homeAddress={"city": "Oxford", "countryOrRegion": "UK"},
        otherAddress={"street": "PO Box 1", "city": "Seattle"},
    )
    assert updated.contact.displayName == "Ada Byron"

    categorized = await graph.set_contact_categories(
        mailbox="shared@example.com",
        contactId="contact-1",
        categories=["Client", "VIP"],
    )
    assert categorized.contact.categories == ["Client", "VIP"]

    added = await graph.add_contact_categories(
        mailbox="shared@example.com",
        contactId="contact-1",
        categories=["Prospect", "Client"],
    )
    assert added.contact.categories == ["Client", "VIP", "Prospect"]

    removed = await graph.remove_contact_categories(
        mailbox="shared@example.com",
        contactId="contact-1",
        categories=["VIP"],
    )
    assert removed.contact.categories == ["Client", "Prospect"]

    cleared = await graph.clear_contact_categories(
        mailbox="shared@example.com",
        contactId="contact-1",
    )
    assert cleared.contact.categories == []

    deleted = await graph.delete_contact(
        mailbox="shared@example.com",
        contactId="contact-1",
        folderId="contacts-folder",
    )
    assert deleted.deleted is True

    assert any("/users/shared@example.com/" in path for _, path, _, _ in requests)
    assert any(
        path.endswith("/contactFolders/contacts-folder/contacts/contact-1")
        and method == "DELETE"
        for method, path, _, _ in requests
    )

    await client.aclose()


@pytest.mark.anyio
async def test_list_and_create_events_and_graph_errors() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content.decode("utf-8")) if request.content else None

        if request.url.path.endswith("/calendarView"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "event-1",
                            "subject": "Planning",
                            "start": {"dateTime": "2026-04-22T16:00:00", "timeZone": "UTC"},
                            "end": {"dateTime": "2026-04-22T17:00:00", "timeZone": "UTC"},
                            "location": {"displayName": "Room 1"},
                            "attendees": [
                                {
                                    "emailAddress": {"address": "attendee@example.com", "name": "Attendee"},
                                    "type": "required",
                                    "status": {"response": "accepted"},
                                }
                            ],
                            "bodyPreview": "Preview",
                            "body": {"contentType": "text", "content": "Agenda"},
                        }
                    ]
                },
            )

        if request.url.path.endswith("/calendar/events"):
            assert body == {
                "subject": "Created event",
                "start": {"dateTime": "2026-04-23T16:00:00", "timeZone": "UTC"},
                "end": {"dateTime": "2026-04-23T17:00:00", "timeZone": "UTC"},
                "attendees": [],
                "body": {"contentType": "HTML", "content": "Agenda"},
            }
            return httpx.Response(
                200,
                json={
                    "id": "event-2",
                    "subject": "Created event",
                    "start": {"dateTime": "2026-04-23T16:00:00", "timeZone": "UTC"},
                    "end": {"dateTime": "2026-04-23T17:00:00", "timeZone": "UTC"},
                    "attendees": [],
                    "bodyPreview": "",
                    "body": {"contentType": "text", "content": ""},
                },
            )

        if request.url.path.endswith("/calendar/events/event-2") and request.method == "PATCH":
            assert body == {
                "subject": "Updated event",
                "start": {"dateTime": "2026-04-23T18:00:00", "timeZone": "UTC"},
                "end": {"dateTime": "2026-04-23T19:00:00", "timeZone": "UTC"},
                "attendees": [
                    {
                        "emailAddress": {"address": "new@example.com"},
                        "type": "required",
                    }
                ],
                "location": {"displayName": "Room 2"},
            }
            return httpx.Response(
                200,
                json={
                    "id": "event-2",
                    "subject": "Updated event",
                    "start": {"dateTime": "2026-04-23T18:00:00", "timeZone": "UTC"},
                    "end": {"dateTime": "2026-04-23T19:00:00", "timeZone": "UTC"},
                    "attendees": [],
                    "bodyPreview": "",
                    "body": {"contentType": "text", "content": ""},
                },
            )

        if request.url.path.endswith("/calendar/events/event-2") and request.method == "DELETE":
            return httpx.Response(204)

        if request.url.path.endswith("/messages/bad-id"):
            return httpx.Response(
                404,
                json={"error": {"code": "ErrorItemNotFound", "message": "No such message"}},
            )

        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    events = await graph.list_events(mailbox="shared@example.com", start="2026-04-22T00:00:00Z", end="2026-04-23T00:00:00Z")
    assert events.events[0].attendees[0].response == "accepted"

    created = await graph.create_event(
        subject="Created event",
        start="2026-04-23T16:00:00",
        end="2026-04-23T17:00:00",
        mailbox="shared@example.com",
        body="Agenda",
        bodyType="html",
    )
    assert created.event.id == "event-2"

    updated = await graph.update_event(
        eventId="event-2",
        mailbox="shared@example.com",
        subject="Updated event",
        start="2026-04-23T18:00:00",
        end="2026-04-23T19:00:00",
        attendees=["new@example.com"],
        location="Room 2",
    )
    assert updated.event.subject == "Updated event"

    deleted = await graph.delete_event(
        eventId="event-2",
        mailbox="shared@example.com",
    )
    assert deleted.deleted is True

    with pytest.raises(RuntimeError, match="ErrorItemNotFound: No such message"):
        await graph.get_message(messageId="bad-id")

    await client.aclose()


@pytest.mark.anyio
async def test_add_and_remove_categories_read_only_categories() -> None:
    patches: list[list[str]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        if request.method == "GET":
            assert request.url.params["$select"] == "categories"
            return httpx.Response(200, json={"id": "msg-1", "categories": ["A", "B"]})
        if request.method == "PATCH":
            body = json.loads(request.content)
            patches.append(body["categories"])
            return httpx.Response(
                200,
                json={
                    "id": "msg-1",
                    "subject": "",
                    "bodyPreview": "",
                    "isDraft": False,
                    "categories": body["categories"],
                },
            )
        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)
    await graph.add_message_categories(messageId="msg-1", categories=["B", "C"])
    await graph.remove_message_categories(messageId="msg-1", categories=["A"])
    assert patches == [["A", "B", "C"], ["B"]]
    await client.aclose()


@pytest.mark.anyio
async def test_get_messages_survives_non_json_errors_and_network_failures(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    sleeps: list[float] = []

    async def fake_sleep(seconds: float) -> None:
        sleeps.append(seconds)

    monkeypatch.setattr(graph_module.anyio, "sleep", fake_sleep)
    calls = {"second": 0, "third": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        ids = [_inner_message_id(inner) for inner in payload["requests"]]
        if "m25" in ids:
            calls["second"] += 1
            if calls["second"] == 1:
                raise httpx.ConnectError("connection reset")
        if "m45" in ids:
            calls["third"] += 1
            return httpx.Response(
                502, text="<html><body>Bad gateway</body></html>"
            )
        return httpx.Response(
            200,
            json={
                "responses": [
                    _message_response(inner, _inner_message_id(inner))
                    for inner in payload["requests"]
                ]
            },
        )

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)
    result = await graph.get_messages(messageIds=[f"m{index}" for index in range(50)])

    # A network error is retried and the chunk then reads normally.
    assert calls["second"] == 2
    assert all(item.message is not None for item in result.messages[:40])
    # An HTML 502 is not JSON: it becomes per-item errors, not a crash.
    assert calls["third"] == 1
    for item in result.messages[40:]:
        assert item.message is None
        assert "502" in item.error
    await client.aclose()


@pytest.mark.anyio
async def test_request_reports_non_json_error_bodies() -> None:
    client = httpx.AsyncClient(
        transport=httpx.MockTransport(
            lambda request: httpx.Response(503, text="<html>Service Unavailable</html>")
        )
    )
    graph = MicrosoftGraphClient(StaticAuthService(), client)
    with pytest.raises(graph_module.GraphRequestError) as raised:
        await graph._request("/me/messages")
    assert raised.value.status == 503
    assert "503" in str(raised.value)
    await client.aclose()


def _xlsx_bytes(
    build,
    *,
    cached: dict[str, str] | None = None,
    dimension: str | None = None,
) -> bytes:
    """Build an .xlsx with openpyxl, then patch the first sheet's XML.

    openpyxl never calculates, so ``cached`` maps a formula's text to the value
    Excel would have saved beside it; ``dimension`` overwrites the recorded
    used range to mimic a writer that left it stale.
    """

    import io
    import re
    import zipfile

    from openpyxl import Workbook

    workbook = Workbook()
    build(workbook)
    raw = io.BytesIO()
    workbook.save(raw)

    patched = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(
        patched, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                for formula, value in (cached or {}).items():
                    # openpyxl writes the empty cached value as <v /> or, when
                    # lxml is installed, as <v></v>.
                    xml = re.sub(
                        rf"<f>{re.escape(formula)}</f><v(?: />|></v>)",
                        f"<f>{formula}</f><v>{value}</v>",
                        xml,
                    )
                if dimension is not None:
                    xml = re.sub(r'<dimension ref="[^"]*" />', f'<dimension ref="{dimension}" />', xml)
                data = xml.encode("utf-8")
            target.writestr(info, data)
    return patched.getvalue()


def _build_deal_model(workbook) -> None:
    from datetime import date

    from openpyxl.workbook.defined_name import DefinedName

    sheet = workbook.active
    sheet.title = "Unit Mix"
    sheet["B2"] = "Units"
    sheet["C2"] = 120
    sheet["B3"] = "Cap Rate"
    sheet["C3"] = 0.055
    sheet["C3"].number_format = "0.0%"
    sheet["B4"] = "Doubled"
    sheet["C4"] = "=C2*2"
    sheet["B5"] = "Never calculated"
    sheet["C5"] = "=C2*3"
    sheet["B6"] = "Close"
    sheet["C6"] = date(2026, 1, 31)

    assumptions = workbook.create_sheet("Assumptions")
    assumptions["A1"] = "Purchase Price"
    assumptions["B1"] = 25_000_000

    hidden = workbook.create_sheet("Calc")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "secret"

    workbook.defined_names["PurchasePrice"] = DefinedName(
        "PurchasePrice", attr_text="Assumptions!$B$1"
    )
    workbook.defined_names["IRR"] = DefinedName(
        "IRR", attr_text="'Unit Mix'!$C$3"
    )
    workbook.defined_names["Broken"] = DefinedName("Broken", attr_text="#REF!")


def _workbook_graph(
    payload: bytes,
    *,
    name: str = "model.xlsx",
    content_type: str = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    odata_type: str = "#microsoft.graph.fileAttachment",
    size: int | None = None,
    downloads: list[int] | None = None,
) -> tuple[httpx.AsyncClient, MicrosoftGraphClient]:
    """Mock Graph: attachment metadata (``$select``, no contentBytes), then the
    raw bytes from ``/$value``. ``downloads`` records each download's size."""

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/messages/msg-1/attachments/xl-1"):
            select = request.url.params.get("$select", "")
            assert select and "contentBytes" not in select
            return httpx.Response(
                200,
                json={
                    "@odata.type": odata_type,
                    "id": "xl-1",
                    "name": name,
                    "contentType": content_type,
                    "size": len(payload) if size is None else size,
                },
            )
        if request.url.path.endswith("/messages/msg-1/attachments/xl-1/$value"):
            if downloads is not None:
                downloads.append(len(payload))
            return httpx.Response(200, content=payload)
        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    return client, MicrosoftGraphClient(StaticAuthService(), client)


@pytest.mark.anyio
async def test_workbook_attachment_outline_lists_sheets_and_names() -> None:
    client, graph = _workbook_graph(_xlsx_bytes(_build_deal_model))

    result = await graph.get_attachment_workbook(messageId="msg-1", attachmentId="xl-1")

    assert result.unsupportedReason is None
    assert result.ranges == []
    assert [
        (sheet.name, sheet.visibility, sheet.dimensions) for sheet in result.sheets
    ] == [
        ("Unit Mix", "visible", "B2:C6"),
        ("Assumptions", "visible", "A1:B1"),
        ("Calc", "hidden", "A1"),
    ]
    assert result.sheets[0].rowCount == 5
    assert result.sheets[0].columnCount == 2
    # The outline counts names but lists none; Broken points at deleted cells.
    assert result.definedNames == []
    assert result.definedNameCount == 2
    assert result.definedNamesSkipped == 1

    named = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", includeDefinedNames=True
    )
    assert {(name.name, name.value) for name in named.definedNames} == {
        ("PurchasePrice", "Assumptions!$B$1"),
        ("IRR", "'Unit Mix'!$C$3"),
    }
    assert len(named.sheets) == 3

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_reads_ranges_names_and_formulas() -> None:
    client, graph = _workbook_graph(
        _xlsx_bytes(_build_deal_model, cached={"C2*2": "240"})
    )

    result = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        ranges=[
            "'unit mix'!B2:C6",
            "PurchasePrice",
            "irr",
            "Calc!A:A",
            "Missing!A1",
            "B2",
        ],
        includeFormulas=True,
        includeNumberFormat=True,
    )

    unit_mix, price, irr, calc, missing, unqualified = result.ranges
    assert result.truncated is False

    assert unit_mix.worksheet == "Unit Mix"
    assert unit_mix.address == "B2:C6"
    assert (unit_mix.rowCount, unit_mix.columnCount) == (5, 2)
    assert unit_mix.values == [
        ["Units", 120],
        ["Cap Rate", 0.055],
        ["Doubled", 240],
        # Excel never saved a result for this formula, so there is none to read.
        ["Never calculated", None],
        ["Close", "2026-01-31T00:00:00"],
    ]
    assert unit_mix.formulas[2] == ["Doubled", "=C2*2"]
    assert unit_mix.formulas[3] == ["Never calculated", "=C2*3"]
    assert unit_mix.numberFormat[1][1] == "0.0%"

    assert (price.worksheet, price.address, price.values) == (
        "Assumptions",
        "B1",
        [[25_000_000]],
    )
    # A bare "IRR" is a defined name, not column IRR.
    assert (irr.worksheet, irr.address, irr.values) == ("Unit Mix", "C3", [[0.055]])
    # Whole columns clamp to the used range, and hidden sheets are readable.
    assert (calc.address, calc.values) == ("A1", [["secret"]])

    assert missing.error == "Workbook has no sheet named 'Missing'"
    assert missing.values is None
    assert "qualify the address with a sheet name" in unqualified.error

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_sheet_mode_and_stale_dimension() -> None:
    def build(workbook) -> None:
        sheet = workbook.active
        sheet.title = "Rent Roll"
        for row in range(1, 6):
            sheet.cell(row=row, column=1, value=f"Unit {row}")
            sheet.cell(row=row, column=3, value=row * 1000)

    # The writer claims only A1 is used; the real data runs to C5.
    client, graph = _workbook_graph(_xlsx_bytes(build, dimension="A1"))

    result = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        sheet="Rent Roll",
        includeLayout=True,
    )

    assert result.sheets[0].dimensions == "A1:C5"
    (whole,) = result.ranges
    assert whole.address == "A1:C5"
    assert whole.values[4] == ["Unit 5", None, 5000]
    assert whole.formulas is None
    assert whole.numberFormat is None

    single = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        ranges=["C2:C3"],
    )
    # One sheet, so an unqualified address needs no sheet name.
    assert single.ranges[0].values == [[2000], [3000]]

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_range_reads_skip_layout_unless_asked() -> None:
    def build(workbook) -> None:
        from openpyxl.workbook.defined_name import DefinedName

        first = workbook.active
        first.title = "Summary"
        first["A1"] = "Price"
        first["B1"] = 1000
        for index in range(40):
            extra = workbook.create_sheet(f"Tab {index}")
            extra["A1"] = index
            workbook.defined_names[f"Input{index}"] = DefinedName(
                f"Input{index}", attr_text=f"'Tab {index}'!$A$1"
            )

    client, graph = _workbook_graph(_xlsx_bytes(build))

    narrow = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["Summary!A1:B1"]
    )
    assert narrow.ranges[0].values == [["Price", 1000]]
    assert narrow.sheets == []
    assert narrow.definedNames == []
    assert len(narrow.model_dump_json()) < 1500

    by_sheet = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", sheet="Summary"
    )
    assert by_sheet.ranges[0].values == [["Price", 1000]]
    assert by_sheet.sheets == [] and by_sheet.definedNames == []

    with_layout = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        ranges=["Summary!A1:B1"],
        includeLayout=True,
    )
    assert len(with_layout.sheets) == 41
    assert with_layout.definedNameCount == 40
    assert with_layout.definedNames == []

    # The discovery call (no ranges, no sheet) still returns the layout.
    layout = await graph.get_attachment_workbook(messageId="msg-1", attachmentId="xl-1")
    assert len(layout.sheets) == 41 and layout.ranges == []

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_names_drop_template_leftovers() -> None:
    def build(workbook) -> None:
        from openpyxl.workbook.defined_name import DefinedName

        sheet = workbook.active
        sheet.title = "Assump"
        sheet["B2"] = 25_000_000
        sheet["B3"] = 0.055
        names = workbook.defined_names
        names["PurchasePrice"] = DefinedName("PurchasePrice", attr_text="Assump!$B$2")
        names["ExitCap"] = DefinedName("ExitCap", attr_text="Assump!$B$3")
        names["TaxRate"] = DefinedName("TaxRate", attr_text="0.012")
        # Leftovers broker templates carry by the hundred.
        for index in range(300):
            junk = "_" * 41 + f"a{index}"
            names[junk] = DefinedName(
                junk, attr_text='{"Assump",#N/A,TRUE,"Proforma";"Assump",#N/A,TRUE,"Proforma"}'
            )
        for index in range(150):
            names[f"OldInput{index}"] = DefinedName(f"OldInput{index}", attr_text="#N/A")
        names["Deleted"] = DefinedName("Deleted", attr_text="Assump!#REF!")
        names["Linked"] = DefinedName("Linked", attr_text="[1]Proforma!$C$4")
        names["RentColumn"] = DefinedName("RentColumn", attr_text="Rents[Rent]")
        names["Hidden"] = DefinedName("Hidden", attr_text="Assump!$B$2", hidden=True)
        names["wrn.Print_All."] = DefinedName(
            "wrn.Print_All.", attr_text='{#N/A,#N/A,FALSE,"Proforma"}'
        )
        sheet.print_area = "A1:B3"

    client, graph = _workbook_graph(_xlsx_bytes(build))

    outline = await graph.get_attachment_workbook(messageId="msg-1", attachmentId="xl-1")
    assert outline.definedNames == []
    assert outline.definedNameCount == 4
    # 300 + 150 leftovers plus Deleted, Linked, Hidden, and wrn.Print_All.;
    # openpyxl lifts the print area onto the sheet, so it is never a name.
    assert outline.definedNamesSkipped == 454
    assert len(outline.model_dump_json()) < 1500

    named = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", includeDefinedNames=True
    )
    assert [name.name for name in named.definedNames] == [
        "PurchasePrice",
        "ExitCap",
        "TaxRate",
        "RentColumn",
    ]
    assert len(named.model_dump_json()) < 2000

    # A name left off the list still resolves when asked for directly.
    with_cells = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        ranges=["Hidden", "PurchasePrice"],
        includeDefinedNames=True,
    )
    assert [r.values for r in with_cells.ranges] == [[[25_000_000]], [[25_000_000]]]
    assert len(with_cells.definedNames) == 4
    assert with_cells.sheets == []

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_cell_budget_truncates() -> None:
    def build(workbook) -> None:
        sheet = workbook.active
        for row in range(1, 11):
            for column in range(1, 4):
                sheet.cell(row=row, column=column, value=row * column)

    client, graph = _workbook_graph(_xlsx_bytes(build))

    result = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        ranges=["A1:C10", "A1:B2"],
        maxCells=10,
    )

    first, second = result.ranges
    assert result.truncated is True
    assert first.truncated is True
    assert first.address == "A1:C3"
    assert first.rowCount == 3
    assert first.values[-1] == [3, 6, 9]
    # One cell of budget is left, which cannot hold a single row of the range.
    assert second.truncated is True
    assert "maxCells=10" in second.error

    await client.aclose()


@pytest.mark.anyio
@pytest.mark.parametrize(
    ("overrides", "reason"),
    [
        ({"name": "legacy.xls", "content_type": "application/vnd.ms-excel"}, "Legacy .xls"),
        ({"name": "binary.xlsb", "content_type": "application/octet-stream"}, "Legacy .xls"),
        ({"name": "notes.txt", "content_type": "text/plain"}, "not an .xlsx"),
        ({"size": 50_000_000}, "maxBytes=10000000"),
        (
            {"odata_type": "#microsoft.graph.itemAttachment"},
            "Only file attachments",
        ),
    ],
)
async def test_workbook_attachment_unsupported(overrides: dict, reason: str) -> None:
    client, graph = _workbook_graph(_xlsx_bytes(_build_deal_model), **overrides)

    result = await graph.get_attachment_workbook(messageId="msg-1", attachmentId="xl-1")

    assert reason in result.unsupportedReason
    assert result.sheets == []

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_rejects_non_zip_payload() -> None:
    # A password-protected workbook is an OLE container, not a zip.
    client, graph = _workbook_graph(b"\xd0\xcf\x11\xe0 not a zip", name="locked.xlsx")

    result = await graph.get_attachment_workbook(messageId="msg-1", attachmentId="xl-1")

    assert "password-protected" in result.unsupportedReason

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_without_openpyxl(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(workbook_reader, "openpyxl", None)
    downloads: list[int] = []
    client, graph = _workbook_graph(_xlsx_bytes(_build_deal_model), downloads=downloads)

    result = await graph.get_attachment_workbook(messageId="msg-1", attachmentId="xl-1")

    assert result.unsupportedReason == "Workbook reading requires the openpyxl package"
    assert downloads == []

    await client.aclose()


@pytest.mark.anyio
async def test_attachment_content_points_workbooks_to_workbook_tool() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "@odata.type": "#microsoft.graph.fileAttachment",
                "id": "xl-1",
                "name": "model.xlsx",
                "contentType": "application/octet-stream",
                "size": 10,
            },
        )

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    result = await graph.get_attachment_content(messageId="msg-1", attachmentId="xl-1")

    assert result.content is None
    assert "mail_get_attachment_workbook" in result.unsupportedReason

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_counts_uncalculated_formulas_in_used_range() -> None:
    def build(workbook) -> None:
        sheet = workbook.active
        sheet["A1"] = "Heading"
        sheet["A2"] = "=1+1"

    # Excel never calculated A2, so it has no cached value.
    client, graph = _workbook_graph(_xlsx_bytes(build))

    result = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        sheet="Sheet",
        includeFormulas=True,
        includeLayout=True,
    )

    assert result.sheets[0].dimensions == "A1:A2"
    (whole,) = result.ranges
    assert whole.address == "A1:A2"
    assert whole.values == [["Heading"], [None]]
    assert whole.formulas == [["Heading"], ["=1+1"]]

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_isolates_unreadable_names() -> None:
    def build(workbook) -> None:
        from openpyxl.workbook.defined_name import DefinedName

        sheet = workbook.active
        sheet["A1"] = "Rent"
        workbook.defined_names["RentColumn"] = DefinedName(
            "RentColumn", attr_text="Rents[Rent]"
        )
        workbook.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="0.012")

    client, graph = _workbook_graph(_xlsx_bytes(build))

    result = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        ranges=["RentColumn", "A1", "TaxRate"],
    )

    assert result.unsupportedReason is None
    structured, cell, constant = result.ranges
    assert "Rents[Rent]" in structured.error
    assert cell.values == [["Rent"]]
    assert "not a single cell range" in constant.error

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_resolves_sheet_qualified_names() -> None:
    def build(workbook) -> None:
        from openpyxl.workbook.defined_name import DefinedName

        _build_deal_model(workbook)
        unit_mix = workbook["Unit Mix"]
        unit_mix.defined_names["Units"] = DefinedName(
            "Units", attr_text="'Unit Mix'!$C$2"
        )

    client, graph = _workbook_graph(_xlsx_bytes(build))

    result = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        ranges=["'Unit Mix'!IRR", "'Unit Mix'!Units", "Assumptions!Units"],
    )

    irr, units, wrong_sheet = result.ranges
    # Not column IRR: the workbook name IRR.
    assert (irr.worksheet, irr.address, irr.values) == ("Unit Mix", "C3", [[0.055]])
    assert (units.address, units.values) == ("C2", [[120]])
    # A sheet-scoped name is only visible from its own sheet.
    assert "not an A1 address or defined name" in wrong_sheet.error

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_pads_ranges_past_the_sheet_end() -> None:
    def build(workbook) -> None:
        sheet = workbook.active
        for row in range(1, 4):
            sheet.cell(row=row, column=1, value=row)

    client, graph = _workbook_graph(_xlsx_bytes(build))

    result = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        ranges=["A2:B5"],
        includeFormulas=True,
        includeNumberFormat=True,
    )

    (padded,) = result.ranges
    assert (padded.address, padded.rowCount, padded.columnCount) == ("A2:B5", 4, 2)
    assert padded.values == [[2, None], [3, None], [None, None], [None, None]]
    assert len(padded.formulas) == 4 and len(padded.numberFormat) == 4

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_stops_download_past_max_bytes() -> None:
    # Graph reports a small size, but the download is larger than maxBytes.
    client, graph = _workbook_graph(b"x" * 5_000, size=100)

    result = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", maxBytes=1_000
    )

    assert result.unsupportedReason == "Workbook size exceeds maxBytes=1000"

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_clamps_limits() -> None:
    client, graph = _workbook_graph(_xlsx_bytes(_build_deal_model), size=50_000_000)

    oversized = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", maxBytes=10**12
    )
    assert oversized.unsupportedReason == "Workbook size exceeds maxBytes=25000000"
    await client.aclose()

    client, graph = _workbook_graph(_xlsx_bytes(_build_deal_model))
    starved = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["'Unit Mix'!B2:C2"], maxCells=0
    )
    assert "maxCells=1 is used up" in starved.ranges[0].error
    await client.aclose()


def _rezip(payload: bytes, replace: dict[str, bytes]) -> bytes:
    import io
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as source, zipfile.ZipFile(
        out, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            target.writestr(info, replace.pop(info.filename, source.read(info.filename)))
        for filename, data in replace.items():
            target.writestr(filename, data)
    return out.getvalue()


@pytest.mark.anyio
async def test_workbook_attachment_refuses_dtds_and_zip_bombs() -> None:
    import io
    import zipfile

    base = _xlsx_bytes(_build_deal_model)
    with zipfile.ZipFile(io.BytesIO(base)) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
    entity = b'<!DOCTYPE worksheet [<!ENTITY a "aaaaaaaa">]>' + sheet_xml

    client, graph = _workbook_graph(_rezip(base, {"xl/worksheets/sheet1.xml": entity}))
    result = await graph.get_attachment_workbook(messageId="msg-1", attachmentId="xl-1")
    assert "document type declaration" in result.unsupportedReason
    await client.aclose()

    # 20 MB of zeros deflates to ~20 KB: far past any real workbook's ratio.
    bomb = _rezip(base, {"xl/media/padding.bin": b"\0" * 20_000_000})
    client, graph = _workbook_graph(bomb)
    result = await graph.get_attachment_workbook(messageId="msg-1", attachmentId="xl-1")
    assert "expands to more than" in result.unsupportedReason
    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_reversed_ranges_cannot_raise_the_cell_budget() -> None:
    # Review repro: a reversed range gave a negative cell count, which raised
    # the remaining budget for the ranges after it.
    def build(workbook) -> None:
        sheet = workbook.active
        for row in range(1, 101):
            sheet.cell(row=row, column=1, value=row)

    client, graph = _workbook_graph(_xlsx_bytes(build))

    result = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["A101:A1", "A1:A100"], maxCells=2
    )
    reversed_range, after = result.ranges
    # Read like Excel does, as A1:A101, within the 2-cell budget.
    assert (reversed_range.address, reversed_range.rowCount) == ("A1:A2", 2)
    assert reversed_range.values == [[1], [2]]
    assert reversed_range.truncated is True
    assert after.values is None
    assert "maxCells=2 is used up" in after.error

    flipped = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["B3:A1"]
    )
    assert flipped.ranges[0].address == "A1:B3"
    assert flipped.ranges[0].values == [[1, None], [2, None], [3, None]]

    outside = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["A1:A1048577", "XFE1", "A1"]
    )
    assert "outside a worksheet" in outside.ranges[0].error
    assert outside.ranges[1].error is not None
    assert outside.ranges[2].values == [[1]]

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_bounds_text_per_cell_and_per_call(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from m365_mcp import workbook_reader
    from m365_mcp.workbook_reader import MAX_WORKBOOK_CELL_CHARS, WORKBOOK_TEXT_CUT_MARKER

    # Scaled down so the test workbook stays under the compression-ratio
    # limit; the per-cell cap is the real one.
    total_chars = 10_000
    monkeypatch.setattr(workbook_reader, "MAX_WORKBOOK_TOTAL_CHARS", total_chars)
    long_text = "x" * (MAX_WORKBOOK_CELL_CHARS * 3)

    def build(workbook) -> None:
        sheet = workbook.active
        for row in range(1, 31):
            sheet.cell(row=row, column=1, value=long_text)
            sheet.cell(row=row, column=2, value=row)

    client, graph = _workbook_graph(_xlsx_bytes(build))

    result = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["A1:B30", "B1"]
    )
    cut, after = result.ranges
    assert cut.values[0][0] == "x" * MAX_WORKBOOK_CELL_CHARS + WORKBOOK_TEXT_CUT_MARKER
    assert cut.values[0][1] == 1
    # Whole rows only, up to the per-call limit; the address matches.
    kept = total_chars // (MAX_WORKBOOK_CELL_CHARS + 1)
    assert cut.rowCount == len(cut.values) == kept
    assert cut.address == f"A1:B{kept}"
    assert cut.truncated is True and result.truncated is True
    assert "character text limit is used up" in after.error

    # One cut cell is enough to flag the range.
    single = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["A1"]
    )
    assert single.ranges[0].truncated is True

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_text_limit_keeps_grids_the_same_height(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from m365_mcp import workbook_reader

    monkeypatch.setattr(workbook_reader, "MAX_WORKBOOK_TOTAL_CHARS", 60)

    def build(workbook) -> None:
        sheet = workbook.active
        for row in range(1, 6):
            sheet.cell(row=row, column=1, value=f"=CONCAT(\"row\",{row})")

    client, graph = _workbook_graph(_xlsx_bytes(build))
    result = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["A1:A5"], includeFormulas=True
    )
    (read,) = result.ranges
    # Values are uncalculated (null), so the formula text is what runs out.
    assert 0 < read.rowCount < 5
    assert len(read.values) == len(read.formulas) == read.rowCount
    assert read.truncated is True

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_caps_ranges_per_call() -> None:
    from m365_mcp.workbook_reader import MAX_WORKBOOK_RANGES

    downloads: list[int] = []
    client, graph = _workbook_graph(_xlsx_bytes(_build_deal_model), downloads=downloads)
    with pytest.raises(ValueError, match=f"at most {MAX_WORKBOOK_RANGES}"):
        await graph.get_attachment_workbook(
            messageId="msg-1",
            attachmentId="xl-1",
            ranges=["'Unit Mix'!C2"] * (MAX_WORKBOOK_RANGES + 1),
        )
    assert downloads == []
    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_extent_scan_is_bounded(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from m365_mcp import workbook_reader

    monkeypatch.setattr(workbook_reader, "WORKBOOK_EXTENT_MAX_CELLS", 5)

    def build(workbook) -> None:
        sheet = workbook.active
        sheet.title = "Big"
        for row in range(1, 11):
            sheet.cell(row=row, column=1, value=row)

    client, graph = _workbook_graph(_xlsx_bytes(build))

    layout = await graph.get_attachment_workbook(messageId="msg-1", attachmentId="xl-1")
    (sheet,) = layout.sheets
    assert sheet.extentIncomplete is True
    assert sheet.dimensions == "A1:A5"

    whole = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", sheet="Big"
    )
    assert whole.ranges[0].truncated is True
    # An explicit address does not depend on the scan.
    explicit = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["A1:A10"]
    )
    assert explicit.ranges[0].truncated is False
    assert explicit.ranges[0].values[-1] == [10]

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_refuses_utf16_dtd() -> None:
    base = _xlsx_bytes(_build_deal_model)
    xml = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE worksheet><worksheet/>'
    client, graph = _workbook_graph(
        _rezip(base, {"xl/worksheets/sheet1.xml": xml.encode("utf-16")})
    )
    result = await graph.get_attachment_workbook(messageId="msg-1", attachmentId="xl-1")
    assert "document type declaration" in result.unsupportedReason
    await client.aclose()


def _build_ten_rows(workbook) -> None:
    from openpyxl.workbook.defined_name import DefinedName

    sheet = workbook.active
    sheet.title = "Assump"
    for row in range(1, 11):
        sheet.cell(row=row, column=1, value=f"Label {row}")
        sheet.cell(row=row, column=2, value=row * 100)
    names = workbook.defined_names
    names["PurchasePrice"] = DefinedName("PurchasePrice", attr_text="Assump!$B$2")
    # What Excel leaves behind when a name's cells are deleted.
    names["Deleted"] = DefinedName("Deleted", attr_text="Assump!#REF!")
    names["Constant"] = DefinedName("Constant", attr_text='"' + "x" * 5_000 + '"')


@pytest.mark.anyio
async def test_workbook_attachment_refuses_deleted_cell_names() -> None:
    client, graph = _workbook_graph(_xlsx_bytes(_build_ten_rows))

    result = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        ranges=["Deleted", "Assump!Deleted", "PurchasePrice"],
    )

    deleted, scoped, price = result.ranges
    for refused in (deleted, scoped):
        assert refused.values is None
        assert "#REF!" in refused.error and "deleted" in refused.error
    # The error is labelled with the sheet the request named.
    assert (deleted.worksheet, deleted.address) == ("", "Deleted")
    assert (scoped.worksheet, scoped.address) == ("Assump", "Deleted")
    assert price.values == [[200]]

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_refuses_row_zero() -> None:
    client, graph = _workbook_graph(_xlsx_bytes(_build_ten_rows))

    result = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        ranges=["A0:B2", "A0", "0:2", "Assump!B0", "A1:B2"],
    )

    *refused, read = result.ranges
    for data in refused:
        assert data.values is None
        assert "rows and columns start at 1" in data.error
    # An error splits sheet and address the same way a successful read does.
    assert (refused[3].worksheet, refused[3].address) == ("Assump", "B0")
    assert refused[0].address == "A0:B2"
    assert read.address == "A1:B2"
    assert read.values == [["Label 1", 100], ["Label 2", 200]]

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_cuts_defined_name_targets_in_errors() -> None:
    client, graph = _workbook_graph(_xlsx_bytes(_build_ten_rows))

    result = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["Constant"] * 20
    )

    for data in result.ranges:
        assert "not a single cell range" in data.error
        assert len(data.error) < 400
    assert len(result.model_dump_json()) < 20 * 500

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_text_limit_covers_names_sheets_and_numbers(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from m365_mcp import workbook_reader

    monkeypatch.setattr(workbook_reader, "MAX_WORKBOOK_TOTAL_CHARS", 1_000)

    def build(workbook) -> None:
        from openpyxl.workbook.defined_name import DefinedName

        first = workbook.active
        first.title = "Numbers"
        for row in range(1, 201):
            first.cell(row=row, column=1, value=1_234_567_890 + row)
        for index in range(40):
            workbook.create_sheet(f"Sheet with a long name {index:03d}")
        for index in range(100):
            name = f"Input_{index:03d}_" + "n" * 30
            workbook.defined_names[name] = DefinedName(name, attr_text="Numbers!$A$1")

    client, graph = _workbook_graph(_xlsx_bytes(build))

    names = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", includeDefinedNames=True
    )
    assert names.definedNameCount == 100
    assert 0 < len(names.definedNames) < 100
    assert names.truncated is True

    layout = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", includeLayout=True, ranges=["Numbers!A1"]
    )
    assert 0 < len(layout.sheets) < 41
    assert layout.truncated is True

    numbers = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["Numbers!A1:A200"]
    )
    (read,) = numbers.ranges
    # Ten-digit numbers: about 1,000 characters' worth, not all 200 rows.
    assert 0 < read.rowCount < 120
    assert read.truncated is True

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_stops_at_the_read_time_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from m365_mcp import workbook_reader

    monkeypatch.setattr(workbook_reader, "WORKBOOK_READ_MAX_SECONDS", -1.0)
    client, graph = _workbook_graph(_xlsx_bytes(_build_ten_rows))

    result = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["A1:B2", "A3:B4"]
    )

    assert result.truncated is True
    for data in result.ranges:
        assert "read time limit is used up" in data.error

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_allows_graph_size_overhead() -> None:
    payload = _xlsx_bytes(_build_ten_rows)
    # Graph's size counts item properties too, so it can exceed the file.
    client, graph = _workbook_graph(payload, size=len(payload) + 2_000)

    result = await graph.get_attachment_workbook(
        messageId="msg-1",
        attachmentId="xl-1",
        ranges=["A1"],
        maxBytes=len(payload) + 10,
    )

    assert result.unsupportedReason is None
    assert result.ranges[0].values == [["Label 1"]]

    await client.aclose()


@pytest.mark.anyio
async def test_workbook_attachment_reads_workbooks_with_svg_images() -> None:
    from tests.test_zip_guard import SVG_WITH_DOCTYPE

    base = _xlsx_bytes(_build_ten_rows)
    client, graph = _workbook_graph(_rezip(base, {"xl/media/image1.svg": SVG_WITH_DOCTYPE}))

    result = await graph.get_attachment_workbook(
        messageId="msg-1", attachmentId="xl-1", ranges=["A1"]
    )

    assert result.unsupportedReason is None
    assert result.ranges[0].values == [["Label 1"]]

    await client.aclose()


def _workbook_with_sheets(count: int) -> bytes:
    import io

    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "S0"
    book.active["A1"] = "aaaa"
    for index in range(1, count):
        book.create_sheet(f"S{index}")["A1"] = "aaaa"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


_READ_DEFAULTS = {
    "sheet": None,
    "includeFormulas": False,
    "includeNumberFormat": False,
    "includeLayout": False,
    "includeDefinedNames": False,
    "maxCells": 1000,
}


def _ticking_clock(
    monkeypatch: pytest.MonkeyPatch, scans: list[str], seconds_per_scan: float
) -> None:
    """A fake clock that only moves when a sheet's extent is measured."""
    now = [0.0]
    real_extent = workbook_reader._worksheet_extent

    def extent(worksheet):
        scans.append(worksheet.title)
        now[0] += seconds_per_scan
        return real_extent(worksheet)

    monkeypatch.setattr(workbook_reader, "monotonic", lambda: now[0])
    monkeypatch.setattr(workbook_reader, "_worksheet_extent", extent)


def test_workbook_layout_stops_measuring_sheets_at_the_time_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    scans: list[str] = []
    _ticking_clock(monkeypatch, scans, 31.0)

    result = workbook_reader.read_workbook(
        _workbook_with_sheets(4), ranges=None, **{**_READ_DEFAULTS, "includeLayout": True}
    )

    # The first scan passes the 30 s limit, so no second one starts.
    assert scans == ["S0"]
    assert [sheet.name for sheet in result.sheets] == ["S0"]
    assert result.truncated is True


def test_workbook_ranges_past_the_time_limit_do_not_measure_sheets(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    scans: list[str] = []
    _ticking_clock(monkeypatch, scans, 20.0)

    result = workbook_reader.read_workbook(
        _workbook_with_sheets(4),
        ranges=["S0!A:A", "S1!A:A", "'S2'!A:A"],
        **_READ_DEFAULTS,
    )

    # S0 is measured and read at 20 s; measuring S1 passes the limit, and
    # S2 is refused without being measured.
    assert scans == ["S0", "S1"]
    first, *later = result.ranges
    assert first.values == [["aaaa"]]
    for data, sheet in zip(later, ["S1", "S2"]):
        assert data.worksheet == sheet
        assert "read time limit is used up" in data.error
    # S2 was never resolved, so it keeps the address as requested.
    assert later[1].address == "A:A"


def test_workbook_text_limit_in_a_later_pass_keeps_the_values(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import io

    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "Sheet"
    for row in range(1, 4):
        book.active.cell(row, 1, "aaaa")
    buffer = io.BytesIO()
    book.save(buffer)
    # "Sheet" (5) + three 4-character values (12) fit; the first number
    # format ("General", 7) does not.
    monkeypatch.setattr(workbook_reader, "MAX_WORKBOOK_TOTAL_CHARS", 20)

    result = workbook_reader.read_workbook(
        buffer.getvalue(),
        ranges=["A1:A3"],
        **{**_READ_DEFAULTS, "includeNumberFormat": True},
    )

    (data,) = result.ranges
    assert data.error is None
    assert data.values == [["aaaa"], ["aaaa"], ["aaaa"]]
    assert data.numberFormat is None
    assert data.truncated is True


@pytest.mark.parametrize("request_text", ["", "  ", "Sheet!", "'Sheet'!$"])
def test_workbook_empty_address_is_an_error(request_text: str) -> None:
    result = workbook_reader.read_workbook(
        _workbook_with_sheets(1), ranges=[request_text], **_READ_DEFAULTS
    )
    (data,) = result.ranges
    assert data.values is None or data.values == []
    assert "names no cells" in (data.error or "")


@pytest.mark.anyio
@pytest.mark.parametrize(
    ("name", "content_type", "read", "reason"),
    [
        ("notes.txt", "text/plain", "content", "Attachment content exceeds maxBytes=1000"),
        ("brief.pdf", "application/pdf", "content", "Attachment content exceeds maxBytes=1000"),
        ("chart.png", "image/png", "image", "Image content exceeds maxBytes=1000"),
        ("deck.pdf", "application/pdf", "pages", "PDF size exceeds maxBytes=1000"),
    ],
)
async def test_attachment_downloads_stop_at_max_bytes(
    name: str, content_type: str, read: str, reason: str
) -> None:
    # Graph reports a tiny size, but the file is 50 chunks of 1,000 bytes.
    metadata_selects: list[str] = []
    chunks_sent: list[int] = []

    class ChunkedFile(httpx.AsyncByteStream):
        async def __aiter__(self):
            for _ in range(50):
                chunks_sent.append(1_000)
                yield b"x" * 1_000

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/messages/msg-1/attachments/att-1/$value"):
            # No Content-Length, so only the running total can stop it.
            return httpx.Response(
                200, stream=ChunkedFile(), headers={"Transfer-Encoding": "chunked"}
            )
        if request.url.path.endswith("/messages/msg-1/attachments/att-1"):
            metadata_selects.append(request.url.params.get("$select", ""))
            return httpx.Response(
                200,
                json={
                    "@odata.type": "#microsoft.graph.fileAttachment",
                    "id": "att-1",
                    "name": name,
                    "contentType": content_type,
                    "size": 10,
                },
            )
        raise AssertionError(f"Unexpected request: {request.method} {request.url}")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)
    call = {
        "content": graph.get_attachment_content,
        "image": graph.get_attachment_image,
        "pages": graph.get_attachment_pdf_pages,
    }[read]

    result = await call(messageId="msg-1", attachmentId="att-1", maxBytes=1_000)

    assert result.unsupportedReason == reason
    # The download stops just past maxBytes instead of reading the whole file.
    assert 0 < len(chunks_sent) <= 2
    # The metadata request names its fields, so it never carries the file.
    assert metadata_selects and metadata_selects[0]
    assert "contentBytes" not in metadata_selects[0]
    # contentId is selected through the fileAttachment cast, so inline images
    # still match the body's cid: references.
    assert "microsoft.graph.fileAttachment/contentId" in metadata_selects[0]

    await client.aclose()


@pytest.mark.anyio
async def test_attachment_download_refuses_oversized_content_length() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/$value"):
            return httpx.Response(200, content=b"x" * 5_000)
        return httpx.Response(
            200,
            json={
                "@odata.type": "#microsoft.graph.fileAttachment",
                "id": "att-1",
                "name": "notes.txt",
                "contentType": "text/plain",
                "size": 10,
            },
        )

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    graph = MicrosoftGraphClient(StaticAuthService(), client)

    result = await graph.get_attachment_content(
        messageId="msg-1", attachmentId="att-1", maxBytes=1_000
    )
    assert result.unsupportedReason == "Attachment content exceeds maxBytes=1000"
    assert result.truncated is True

    await client.aclose()
